"""Tests unitarios para los reporters de SchemaDiff.

Verifican que cada reporter produce un fichero con el contenido esperado
usando datos sintéticos (sin conexión a BD).
"""

from __future__ import annotations

import zipfile

from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import pytest

from util_bbdd.comparar.models import (
    ColumnDiff,
    CommentDiff,
    ForeignKeyDiff,
    IndexDiff,
    RoutineDiff,
    SchemaDiff,
    TableDiff,
    ViewDiff,
)
from util_bbdd.comparar.reporters import (
    ExcelDiffReporter,
    HtmlDiffReporter,
    MarkdownDiffReporter,
    SchemaDiffReporter,
    SqlDiffReporter,
    WordDiffReporter,
)
from util_bbdd.comparar.reporters._base import (
    DdlNormalizer,
    build_diff_summary,
    build_output_stem,
    escape_mermaid_label,
    sanitize_filename,
)

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture()
def empty_diff() -> SchemaDiff:
    """SchemaDiff sin diferencias."""
    return SchemaDiff(source_label="DB_A", target_label="DB_B")


@pytest.fixture()
def rich_diff() -> SchemaDiff:
    """SchemaDiff con un ejemplo completo de cada tipo de diferencia."""
    return SchemaDiff(
        source_label="prod",
        target_label="staging",
        summary="Comparación de entornos",
        table_diffs=(
            TableDiff(
                table_name="users",
                status="modified",
                column_diffs=(
                    ColumnDiff(
                        column_name="email",
                        attribute="type",
                        source_value="VARCHAR(100)",
                        target_value="VARCHAR(255)",
                    ),
                    ColumnDiff(
                        column_name="deleted_at",
                        attribute="existence",
                        source_value="present",
                        target_value="absent",
                    ),
                ),
                index_diffs=(
                    IndexDiff(
                        index_name="idx_email",
                        table_name="users",
                        status="modified",
                        source_columns="email",
                        target_columns="email, created_at",
                        source_unique=True,
                        target_unique=False,
                    ),
                ),
                fk_diffs=(
                    ForeignKeyDiff(
                        fk_name="fk_users_company",
                        table_name="users",
                        status="added",
                        target_columns="company_id",
                        target_referred_table="companies",
                    ),
                ),
                comment_diffs=(
                    CommentDiff(
                        object_name="users",
                        column_name="email",
                        source_comment="Email del usuario",
                        target_comment="Dirección de correo",
                    ),
                ),
            ),
            TableDiff(table_name="orders_archive", status="added"),
            TableDiff(table_name="legacy_log", status="removed"),
        ),
        view_diffs=(
            ViewDiff(
                view_name="v_active_users",
                status="modified",
                source_definition="SELECT * FROM users WHERE active=1",
                target_definition="SELECT id, name FROM users WHERE active=1",
            ),
        ),
        routine_diffs=(
            RoutineDiff(
                routine_name="sp_cleanup",
                routine_type="PROCEDURE",
                status="removed",
                source_definition="BEGIN DELETE FROM legacy_log; END",
            ),
        ),
    )


# ---------------------------------------------------------------------------
# Tests de utilidades base
# ---------------------------------------------------------------------------


class TestBase:
    def test_sanitize_filename_special_chars(self) -> None:
        assert sanitize_filename("DB/prod:2026") == "DB_prod_2026"

    def test_sanitize_filename_truncation(self) -> None:
        long_name = "a" * 100
        assert len(sanitize_filename(long_name)) == 64

    def test_build_output_stem(self, empty_diff: SchemaDiff) -> None:
        assert build_output_stem(empty_diff) == "DB_A_vs_DB_B_diff"

    def test_escape_mermaid_label(self) -> None:
        assert '"' not in escape_mermaid_label('table"name')
        assert "[" not in escape_mermaid_label("table[name]")

    def test_build_diff_summary_empty(self, empty_diff: SchemaDiff) -> None:
        s = build_diff_summary(empty_diff)
        assert s.tables_added == 0
        assert s.total_tables == 0
        assert not s.has_differences

    def test_build_diff_summary_rich(self, rich_diff: SchemaDiff) -> None:
        s = build_diff_summary(rich_diff)
        assert s.tables_added == 1
        assert s.tables_removed == 1
        assert s.tables_modified == 1
        assert s.total_views == 1
        assert s.total_routines == 1
        assert s.has_differences


# ---------------------------------------------------------------------------
# Tests de MarkdownDiffReporter
# ---------------------------------------------------------------------------


class TestMarkdownDiffReporter:
    def test_render_empty_diff(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.md"
        result = MarkdownDiffReporter().render(empty_diff, output)
        assert result == output
        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "DB_A" in content
        assert "DB_B" in content
        assert "Sin diferencias" in content

    def test_render_rich_diff(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.md"
        MarkdownDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # Diagrama Mermaid presente
        assert "```mermaid" in content
        assert "graph LR" in content
        # Tablas
        assert "users" in content
        assert "orders_archive" in content
        assert "legacy_log" in content
        # Columnas
        assert "email" in content
        assert "VARCHAR(100)" in content
        # Vistas
        assert "v_active_users" in content
        # Rutinas
        assert "sp_cleanup" in content

    def test_render_creates_parent_dir(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "nested" / "dir" / "diff.md"
        MarkdownDiffReporter().render(empty_diff, output)
        assert output.exists()

    def test_summary_line_in_markdown(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.md"
        MarkdownDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert "Comparación de entornos" in content


# ---------------------------------------------------------------------------
# Tests de HtmlDiffReporter
# ---------------------------------------------------------------------------


class TestHtmlDiffReporter:
    def test_render_empty_diff(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.html"
        result = HtmlDiffReporter().render(empty_diff, output)
        assert result == output
        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "<!DOCTYPE html>" in content
        assert "DB_A" in content
        assert "DB_B" in content

    def test_render_rich_diff(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.html"
        HtmlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # Contiene el diagrama Mermaid
        assert "mermaid" in content
        # Tablas presentes
        assert "users" in content
        assert "orders_archive" in content
        # DDL de vista escapado
        assert "v_active_users" in content
        # Script Mermaid CDN
        assert "cdn.jsdelivr.net" in content

    def test_filter_toolbar_present_when_diffs_exist(
        self, tmp_path: Path, rich_diff: SchemaDiff
    ) -> None:
        """El toolbar de filtros aparece con los 5 botones cuando hay diferencias."""
        output = tmp_path / "diff.html"
        HtmlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert "filter-btn" in content
        assert 'data-filter="tables-added"' in content
        assert 'data-filter="tables-removed"' in content
        assert 'data-filter="tables-modified"' in content
        assert 'data-filter="views"' in content
        assert 'data-filter="routines"' in content

    def test_filter_data_cat_attributes(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        """Cada bloque tiene el atributo data-cat correcto según su tipo."""
        output = tmp_path / "diff.html"
        HtmlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert 'data-cat="tables-added"' in content
        assert 'data-cat="tables-removed"' in content
        assert 'data-cat="tables-modified"' in content
        assert 'data-cat="views"' in content
        assert 'data-cat="routines"' in content

    def test_filter_toolbar_absent_for_empty_diff(
        self, tmp_path: Path, empty_diff: SchemaDiff
    ) -> None:
        """El toolbar NO aparece cuando no hay diferencias (ninguna categoría tiene datos)."""
        output = tmp_path / "diff.html"
        HtmlDiffReporter().render(empty_diff, output)
        content = output.read_text(encoding="utf-8")
        # No debe haber la barra de controles en el DOM cuando no hay diffs
        assert '<div class="controls-bar">' not in content

    def test_html_search_input_present(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        """El HTML contiene un input de búsqueda de texto cuando hay diferencias."""
        output = tmp_path / "diff.html"
        HtmlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert 'type="search"' in content
        assert 'id="search-input"' in content

    def test_ddl_is_html_escaped(self, tmp_path: Path) -> None:
        """Verifica que el DDL con caracteres peligrosos se escapa en HTML."""
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            view_diffs=(
                ViewDiff(
                    view_name="v_xss",
                    status="modified",
                    source_definition="<script>alert(1)</script>",
                ),
            ),
        )
        output = tmp_path / "diff.html"
        HtmlDiffReporter().render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "<script>alert(1)</script>" not in content
        assert "&lt;script&gt;" in content


# ---------------------------------------------------------------------------
# Tests de ExcelDiffReporter
# ---------------------------------------------------------------------------


class TestExcelDiffReporter:
    def test_render_empty_diff(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.xlsx"
        result = ExcelDiffReporter().render(empty_diff, output)
        assert result == output
        assert output.exists()
        # Es un ZIP válido (xlsx)
        assert zipfile.is_zipfile(output)

    def test_render_rich_diff_sheets(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        import openpyxl

        output = tmp_path / "diff.xlsx"
        ExcelDiffReporter().render(rich_diff, output)
        wb = openpyxl.load_workbook(output)
        sheet_names = wb.sheetnames
        assert "Resumen" in sheet_names
        assert "Tablas" in sheet_names
        assert "Columnas" in sheet_names
        assert "Indices" in sheet_names
        assert "ForeignKeys" in sheet_names
        assert "Vistas" in sheet_names
        assert "Rutinas" in sheet_names

    def test_summary_sheet_content(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        import openpyxl

        output = tmp_path / "diff.xlsx"
        ExcelDiffReporter().render(rich_diff, output)
        wb = openpyxl.load_workbook(output)
        ws = wb["Resumen"]
        values = [str(row[1].value or "") for row in ws.iter_rows(min_row=2)]
        assert "prod" in values
        assert "staging" in values

    def test_render_empty_diff_no_extra_sheets(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        import openpyxl

        output = tmp_path / "diff.xlsx"
        ExcelDiffReporter().render(empty_diff, output)
        wb = openpyxl.load_workbook(output)
        # Solo Resumen cuando no hay diferencias
        assert wb.sheetnames == ["Resumen"]


# ---------------------------------------------------------------------------
# Tests de WordDiffReporter
# ---------------------------------------------------------------------------


class TestWordDiffReporter:
    def test_render_empty_diff(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.docx"
        result = WordDiffReporter().render(empty_diff, output)
        assert result == output
        assert output.exists()
        # docx es un ZIP válido
        assert zipfile.is_zipfile(output)

    def test_render_rich_diff_content(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        from docx import Document

        output = tmp_path / "diff.docx"
        WordDiffReporter().render(rich_diff, output)
        doc = Document(str(output))
        # Labels prod/staging están en celdas de tabla (resumen), no en párrafos
        para_text = " ".join(p.text for p in doc.paragraphs)
        table_text = " ".join(
            cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells
        )
        all_text = para_text + " " + table_text
        assert "prod" in all_text
        assert "staging" in all_text
        assert "users" in para_text
        assert "v_active_users" in para_text
        assert "sp_cleanup" in para_text

    def test_template_report_content_is_present(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """Con plantilla vacía, el contenido del reporte se genera correctamente."""
        from docx import Document

        # Plantilla vacía (solo estilos, sin contenido de ejemplo)
        tpl_path = tmp_path / "template.docx"
        Document().save(str(tpl_path))

        output = tmp_path / "diff_from_tpl.docx"
        WordDiffReporter(template_path=tpl_path).render(empty_diff, output)

        doc = Document(str(output))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "Reporte de diferencias" in all_text

    def test_template_styles_are_inherited(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """Usando plantilla el documento se genera correctamente (es un ZIP válido)."""
        from docx import Document

        tpl_path = tmp_path / "template.docx"
        Document().save(str(tpl_path))

        output = tmp_path / "diff_from_tpl.docx"
        WordDiffReporter(template_path=tpl_path).render(empty_diff, output)
        assert zipfile.is_zipfile(output)

    def test_render_raises_on_directory_output(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """Pasar un directorio como output lanza ValueError con mensaje claro."""
        with pytest.raises(ValueError, match="ruta completa del fichero"):
            WordDiffReporter().render(empty_diff, output=tmp_path)

    def test_render_raises_on_path_without_extension(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """Pasar una ruta sin extensión lanza ValueError."""
        with pytest.raises(ValueError, match="ruta completa del fichero"):
            WordDiffReporter().render(empty_diff, output=tmp_path / "nodotx")

    def test_word_heading1_sections(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        """Las 5 categorías usan estilo Heading 1 en el documento Word."""
        from docx import Document

        output = tmp_path / "diff.docx"
        WordDiffReporter().render(rich_diff, output)
        doc = Document(str(output))

        heading1_texts = [
            p.text for p in doc.paragraphs
            if p.style.name == "Heading 1"
        ]
        # rich_diff tiene tablas added+removed+modified, vistas y rutinas
        assert "Tablas añadidas" in heading1_texts
        assert "Tablas eliminadas" in heading1_texts
        assert "Tablas modificadas" in heading1_texts
        assert "Vistas con diff" in heading1_texts
        assert "Rutinas con diff" in heading1_texts


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter
# ---------------------------------------------------------------------------


class TestSqlDiffReporter:
    def test_render_empty_diff(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        result = SqlDiffReporter().render(empty_diff, output)
        assert result == output
        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "Sin diferencias" in content

    def test_render_rich_diff_contains_sections(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # Secciones presentes
        assert "TABLAS ELIMINADAS" in content
        assert "TABLAS AÑADIDAS" in content
        assert "TABLAS MODIFICADAS" in content
        assert "VISTAS" in content
        assert "RUTINAS" in content

    def test_drop_table_for_removed(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # legacy_log tiene status="removed"
        assert "DROP TABLE" in content
        assert "legacy_log" in content

    def test_create_placeholder_for_added(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # orders_archive tiene status="added" — sin DDL completo
        assert "orders_archive" in content
        assert "⚠" in content  # advertencia presente

    def test_alter_table_for_modified_columns(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER TABLE" in content
        assert "users" in content
        # email column type diff presente
        assert "email" in content
        assert "VARCHAR(100)" in content

    def test_postgresql_dialect_uses_alter_column_type(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER COLUMN" in content
        assert "TYPE" in content

    def test_header_contains_dialect(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle").render(empty_diff, output)
        content = output.read_text(encoding="utf-8")
        assert "oracle" in content

    def test_view_ddl_in_script(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # v_active_users tiene status="modified" con source_definition
        assert "v_active_users" in content

    def test_routine_drop_in_script(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # sp_cleanup tiene status="removed"
        assert "sp_cleanup" in content
        assert "DROP" in content

    def test_invalid_dialect_raises(self, tmp_path: Path) -> None:
        with pytest.raises(ValueError, match="no soportado"):
            SqlDiffReporter(dialect="cassandra")

    def test_drop_fk_before_drop_table(self, tmp_path: Path) -> None:
        """FKs de tablas eliminadas se dropean antes que la tabla."""
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="orders",
                    status="removed",
                    fk_diffs=(
                        ForeignKeyDiff(
                            fk_name="fk_orders_user",
                            table_name="orders",
                            status="removed",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        # Usar cadenas específicas de las instrucciones DDL, no del título de sección
        drop_fk_pos    = content.index("DROP FOREIGN KEY fk_orders_user")
        drop_table_pos = content.index("DROP TABLE IF EXISTS orders")
        assert drop_fk_pos < drop_table_pos

    def test_index_create_for_added_index(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # idx_email tiene status="modified" → DROP + CREATE
        assert "idx_email" in content
        assert "CREATE" in content and "INDEX" in content

    def test_add_constraint_for_added_fk(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        # fk_users_company tiene status="added" → ADD CONSTRAINT
        assert "fk_users_company" in content
        assert "ADD CONSTRAINT" in content


# ---------------------------------------------------------------------------
# Tests de SchemaDiffReporter (facade)
# ---------------------------------------------------------------------------


class TestSchemaDiffReporter:
    def test_render_html(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(empty_diff, output_dir=tmp_path, format="html")
        assert path.suffix == ".html"
        assert path.exists()

    def test_render_excel(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(empty_diff, output_dir=tmp_path, format="excel")
        assert path.suffix == ".xlsx"
        assert path.exists()

    def test_render_word(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(empty_diff, output_dir=tmp_path, format="word")
        assert path.suffix == ".docx"
        assert path.exists()

    def test_render_markdown(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(empty_diff, output_dir=tmp_path, format="markdown")
        assert path.suffix == ".md"
        assert path.exists()

    def test_render_invalid_format(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        with pytest.raises(ValueError, match="no soportado"):
            SchemaDiffReporter().render(empty_diff, output_dir=tmp_path, format="pdf")

    def test_render_all_returns_four_files(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """render_all con subset explícito de 4 formatos sin sql."""
        paths = SchemaDiffReporter().render_all(
            empty_diff, output_dir=tmp_path,
            formats=["html", "excel", "word", "markdown"],
        )
        assert len(paths) == 4
        extensions = {p.suffix for p in paths}
        assert extensions == {".html", ".xlsx", ".docx", ".md"}

    def test_render_all_subset(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        paths = SchemaDiffReporter().render_all(
            empty_diff, output_dir=tmp_path, formats=["html", "markdown"]
        )
        assert len(paths) == 2

    def test_render_custom_filename(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(
            empty_diff, output_dir=tmp_path, format="html", filename="custom.html"
        )
        assert path.name == "custom.html"

    def test_render_creates_output_dir(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        new_dir = tmp_path / "nested" / "output"
        SchemaDiffReporter().render(empty_diff, output_dir=new_dir, format="markdown")
        assert new_dir.exists()

    def test_auto_filename_derived_from_labels(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(rich_diff, output_dir=tmp_path, format="markdown")
        assert "prod" in path.stem
        assert "staging" in path.stem

    def test_word_template_path_passed_to_reporter(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """word_template_path se propaga al WordDiffReporter desde la facade."""
        from docx import Document

        tpl_path = tmp_path / "corp.docx"
        Document().save(str(tpl_path))

        path = SchemaDiffReporter(word_template_path=tpl_path).render(
            empty_diff, output_dir=tmp_path, format="word"
        )
        assert path.suffix == ".docx"
        assert zipfile.is_zipfile(path)

    def test_word_template_not_found_raises(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """Si la plantilla no existe la facade propaga FileNotFoundError."""
        missing = tmp_path / "no_existe.docx"
        reporter = SchemaDiffReporter(word_template_path=missing)
        with pytest.raises(FileNotFoundError):
            reporter.render(empty_diff, output_dir=tmp_path, format="word")

    def test_render_sql(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter().render(empty_diff, output_dir=tmp_path, format="sql")
        assert path.suffix == ".sql"
        assert path.exists()

    def test_render_sql_with_dialect(self, tmp_path: Path, rich_diff: SchemaDiff) -> None:
        path = SchemaDiffReporter(sql_dialect="postgresql").render(
            rich_diff, output_dir=tmp_path, format="sql"
        )
        assert path.suffix == ".sql"
        content = path.read_text(encoding="utf-8")
        assert "postgresql" in content.lower()

    def test_render_all_returns_five_files(self, tmp_path: Path, empty_diff: SchemaDiff) -> None:
        """render_all sin filtro genera los 5 formatos incluyendo sql."""
        paths = SchemaDiffReporter().render_all(empty_diff, output_dir=tmp_path)
        assert len(paths) == 5
        extensions = {p.suffix for p in paths}
        assert extensions == {".html", ".xlsx", ".docx", ".md", ".sql"}


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter con source_inspector
# ---------------------------------------------------------------------------


class TestSqlDiffReporterWithInspector:
    """Verifica la generación de SQL completo cuando se proporciona source_inspector."""

    @staticmethod
    def _make_inspector(
        columns: list[dict[str, Any]],
        pk_cols: list[str],
    ) -> MagicMock:
        """Crea un Inspector mock con ``get_columns`` y ``get_pk_constraint`` configurados."""
        mock: MagicMock = MagicMock()
        mock.get_columns.return_value = columns
        mock.get_pk_constraint.return_value = {
            "constrained_columns": pk_cols,
            "name": None,
        }
        return mock

    # ------------------------------------------------------------------
    # CREATE TABLE completo para tablas added
    # ------------------------------------------------------------------

    def test_create_table_full_ddl_with_inspector(self, tmp_path: Path) -> None:
        """Con Inspector, genera CREATE TABLE con columnas reales en lugar de placeholder."""
        cols: list[dict[str, Any]] = [
            {"name": "id",   "type": "INTEGER",      "nullable": False, "default": None},
            {"name": "name", "type": "VARCHAR(100)",  "nullable": True,  "default": None},
        ]
        inspector = self._make_inspector(cols, ["id"])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="orders_new", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")

        assert "CREATE TABLE" in content
        assert "orders_new" in content
        assert "id INTEGER" in content
        assert "name VARCHAR(100)" in content
        assert "PRIMARY KEY (id)" in content
        # No se emite el placeholder ⚠
        assert "⚠ TABLA NUEVA" not in content

    def test_create_table_if_not_exists_for_postgresql(self, tmp_path: Path) -> None:
        """Dialecto postgresql usa CREATE TABLE IF NOT EXISTS."""
        cols: list[dict[str, Any]] = [
            {"name": "id", "type": "INTEGER", "nullable": False, "default": None},
        ]
        inspector = self._make_inspector(cols, ["id"])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="t", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql", source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CREATE TABLE IF NOT EXISTS t" in content

    def test_create_table_generic_dialect_no_if_not_exists(self, tmp_path: Path) -> None:
        """Dialecto generic usa CREATE TABLE sin IF NOT EXISTS."""
        cols: list[dict[str, Any]] = [
            {"name": "id", "type": "INTEGER", "nullable": False, "default": None},
        ]
        inspector = self._make_inspector(cols, ["id"])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="t", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="generic", source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CREATE TABLE t (" in content
        assert "IF NOT EXISTS" not in content

    def test_create_table_fallback_placeholder_when_inspector_returns_empty(
        self, tmp_path: Path
    ) -> None:
        """Si el Inspector devuelve lista vacía, se emite el placeholder ⚠."""
        inspector = self._make_inspector([], [])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="ghost", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "⚠ TABLA NUEVA" in content
        assert "ghost" in content

    def test_create_table_fallback_when_inspector_raises(self, tmp_path: Path) -> None:
        """Si el Inspector lanza excepción en get_columns, se usa placeholder."""
        inspector = MagicMock()
        inspector.get_columns.side_effect = Exception("connection lost")
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="broken", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "⚠ TABLA NUEVA" in content

    def test_create_table_not_null_column_in_ddl(self, tmp_path: Path) -> None:
        """Una columna NOT NULL sin DEFAULT aparece con NOT NULL en CREATE TABLE."""
        cols: list[dict[str, Any]] = [
            {"name": "code", "type": "VARCHAR(10)", "nullable": False, "default": None},
        ]
        inspector = self._make_inspector(cols, ["code"])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="products", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "code VARCHAR(10) NOT NULL" in content

    def test_create_table_column_with_default(self, tmp_path: Path) -> None:
        """Una columna con DEFAULT lo incluye en la definición."""
        cols: list[dict[str, Any]] = [
            {"name": "status", "type": "VARCHAR(20)", "nullable": True, "default": "'active'"},
        ]
        inspector = self._make_inspector(cols, [])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="accounts", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DEFAULT 'active'" in content

    # ------------------------------------------------------------------
    # ADD COLUMN con tipo real para columnas existence-added
    # ------------------------------------------------------------------

    def test_add_column_with_real_type_from_inspector(self, tmp_path: Path) -> None:
        """Con Inspector, ADD COLUMN incluye el tipo real en lugar de <TIPO_ORIGEN>."""
        cols: list[dict[str, Any]] = [
            {"name": "id",    "type": "INTEGER",      "nullable": False, "default": None},
            {"name": "email", "type": "VARCHAR(200)", "nullable": True,  "default": None},
        ]
        inspector = self._make_inspector(cols, ["id"])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="users",
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name="email",
                            attribute="existence",
                            source_value="present",
                            target_value="absent",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ADD COLUMN email VARCHAR(200)" in content
        assert "<TIPO_ORIGEN>" not in content

    def test_add_column_uses_exists_value_from_real_comparator(self, tmp_path: Path) -> None:
        """Con valor 'exists' (semántica real del SchemaComparator), genera ADD COLUMN."""
        cols: list[dict[str, Any]] = [
            {"name": "phone", "type": "VARCHAR(20)", "nullable": True, "default": None},
        ]
        inspector = self._make_inspector(cols, [])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="contacts",
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name="phone",
                            attribute="existence",
                            source_value="exists",   # valor real del comparador
                            target_value="missing",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ADD COLUMN phone VARCHAR(20)" in content

    def test_add_column_not_null_without_default_emits_warning(self, tmp_path: Path) -> None:
        """ADD COLUMN NOT NULL sin DEFAULT emite comentario de advertencia."""
        cols: list[dict[str, Any]] = [
            {"name": "code", "type": "VARCHAR(10)", "nullable": False, "default": None},
        ]
        inspector = self._make_inspector(cols, [])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="products",
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name="code",
                            attribute="existence",
                            source_value="present",
                            target_value="absent",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "NOT NULL" in content
        assert "⚠" in content  # advertencia de NOT NULL sin DEFAULT

    def test_add_column_fallback_placeholder_without_inspector(
        self, tmp_path: Path, rich_diff: SchemaDiff
    ) -> None:
        """Sin Inspector, ADD COLUMN emite placeholder <TIPO_ORIGEN> (comportamiento original)."""
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(rich_diff, output)
        content = output.read_text(encoding="utf-8")
        assert "<TIPO_ORIGEN>" in content

    # ------------------------------------------------------------------
    # DROP COLUMN (no requiere inspector — comportamiento existente)
    # ------------------------------------------------------------------

    def test_drop_column_for_existence_missing_source(self, tmp_path: Path) -> None:
        """Columna ausente en source (present en target) genera DROP COLUMN."""
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="orders",
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name="legacy_col",
                            attribute="existence",
                            source_value="absent",
                            target_value="present",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DROP COLUMN legacy_col" in content

    def test_drop_column_uses_missing_value_from_real_comparator(self, tmp_path: Path) -> None:
        """Con valor 'missing' (semántica real del SchemaComparator), genera DROP COLUMN."""
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="orders",
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name="old_col",
                            attribute="existence",
                            source_value="missing",  # valor real del comparador
                            target_value="exists",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter().render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DROP COLUMN old_col" in content

    # ------------------------------------------------------------------
    # Dialecto SQL Server (ADD sin COLUMN)
    # ------------------------------------------------------------------

    def test_sqlserver_add_column_no_column_keyword(self, tmp_path: Path) -> None:
        """SQL Server usa ALTER TABLE … ADD {col} {type} sin la palabra COLUMN."""
        cols: list[dict[str, Any]] = [
            {"name": "status", "type": "NVARCHAR(50)", "nullable": True, "default": None},
        ]
        inspector = self._make_inspector(cols, [])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name="orders",
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name="status",
                            attribute="existence",
                            source_value="present",
                            target_value="absent",
                        ),
                    ),
                ),
            ),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="sqlserver", source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ADD status NVARCHAR(50)" in content
        assert "ADD COLUMN status" not in content

    # ------------------------------------------------------------------
    # Facade: propagación del Inspector
    # ------------------------------------------------------------------

    def test_facade_passes_inspector_to_sql_reporter(self, tmp_path: Path) -> None:
        """SchemaDiffReporter propaga sql_source_inspector al SqlDiffReporter."""
        cols: list[dict[str, Any]] = [
            {"name": "id", "type": "INTEGER", "nullable": False, "default": None},
        ]
        inspector = self._make_inspector(cols, ["id"])
        diff = SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(TableDiff(table_name="new_t", status="added"),),
        )
        path = SchemaDiffReporter(sql_source_inspector=inspector).render(
            diff, output_dir=tmp_path, format="sql"
        )
        content = path.read_text(encoding="utf-8")
        assert "CREATE TABLE" in content
        assert "id INTEGER" in content
        assert "⚠ TABLA NUEVA" not in content

    def test_facade_without_inspector_preserves_placeholder(
        self, tmp_path: Path, rich_diff: SchemaDiff
    ) -> None:
        """Sin Inspector en la facade, el SQL reporter emite placeholders."""
        path = SchemaDiffReporter().render(rich_diff, output_dir=tmp_path, format="sql")
        content = path.read_text(encoding="utf-8")
        assert "⚠ TABLA NUEVA" in content  # orders_archive es added sin DDL


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter — CommentDiff rendering
# ---------------------------------------------------------------------------


class TestSqlDiffReporterCommentDiffs:
    """Verifica que CommentDiff se renderiza con la sintaxis correcta por dialecto."""

    @staticmethod
    def _diff_with_comment(
        table: str,
        column_name: str | None,
        source_comment: str,
    ) -> SchemaDiff:
        return SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name=table,
                    status="modified",
                    comment_diffs=(
                        CommentDiff(
                            object_name=table,
                            column_name=column_name,
                            source_comment=source_comment,
                            target_comment="",
                        ),
                    ),
                ),
            ),
        )

    def test_table_comment_postgresql(self, tmp_path: Path) -> None:
        """PostgreSQL tabla → COMMENT ON TABLE."""
        diff = self._diff_with_comment("users", None, "Tabla de usuarios")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "COMMENT ON TABLE users IS 'Tabla de usuarios'" in content

    def test_column_comment_postgresql(self, tmp_path: Path) -> None:
        """PostgreSQL columna → COMMENT ON COLUMN."""
        diff = self._diff_with_comment("users", "email", "Correo del usuario")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "COMMENT ON COLUMN users.email IS 'Correo del usuario'" in content

    def test_table_comment_oracle(self, tmp_path: Path) -> None:
        """Oracle tabla → COMMENT ON TABLE."""
        diff = self._diff_with_comment("orders", None, "Pedidos")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "COMMENT ON TABLE orders IS 'Pedidos'" in content

    def test_table_comment_mariadb(self, tmp_path: Path) -> None:
        """MariaDB tabla → ALTER TABLE … COMMENT."""
        diff = self._diff_with_comment("products", None, "Catálogo de productos")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER TABLE products COMMENT = 'Cat" in content

    def test_column_comment_mariadb_without_inspector(self, tmp_path: Path) -> None:
        """MariaDB columna sin inspector → MODIFY COLUMN con <TIPO>."""
        diff = self._diff_with_comment("users", "email", "Correo")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN email <TIPO> COMMENT 'Correo'" in content

    def test_column_comment_mariadb_with_inspector(self, tmp_path: Path) -> None:
        """MariaDB columna con inspector → MODIFY COLUMN con tipo real."""
        from unittest.mock import MagicMock

        mock = MagicMock()
        mock.get_columns.return_value = [
            {"name": "email", "type": "VARCHAR(200)", "nullable": True, "default": None}
        ]
        diff = self._diff_with_comment("users", "email", "Email")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb", source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN email VARCHAR(200) COMMENT 'Email'" in content

    def test_sqlite_comment_placeholder(self, tmp_path: Path) -> None:
        """SQLite → placeholder ⚠ para comentarios."""
        diff = self._diff_with_comment("items", None, "Items")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="sqlite").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "⚠" in content
        assert "SQLite" in content

    def test_sqlserver_comment_placeholder(self, tmp_path: Path) -> None:
        """SQL Server → placeholder ⚠ para comentarios."""
        diff = self._diff_with_comment("items", "name", "Nombre del item")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="sqlserver").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "⚠" in content
        assert "sp_addextendedproperty" in content

    def test_comment_with_single_quote_escaped(self, tmp_path: Path) -> None:
        """Comillas simples en comentarios se escapan correctamente."""
        diff = self._diff_with_comment("users", None, "L'utilisateur")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "L''utilisateur" in content

    def test_comment_section_appears_in_modified_tables(self, tmp_path: Path) -> None:
        """La sección 'Comentarios:' aparece en tablas modificadas con comment_diffs."""
        diff = self._diff_with_comment("users", None, "Comentario nuevo")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="generic").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "Comentarios:" in content


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter — CREATE TABLE con FKs e índices desde Inspector
# ---------------------------------------------------------------------------


class TestSqlDiffReporterCreateTableComplete:
    """Verifica CREATE TABLE completo: columnas + PK + FK + índices secundarios."""

    @staticmethod
    def _make_full_inspector(
        columns: list[dict[str, Any]],
        pk_cols: list[str],
        foreign_keys: list[dict[str, Any]] | None = None,
        indexes: list[dict[str, Any]] | None = None,
    ) -> MagicMock:
        mock = MagicMock()
        mock.get_columns.return_value = columns
        mock.get_pk_constraint.return_value = {"constrained_columns": pk_cols, "name": None}
        mock.get_foreign_keys.return_value = foreign_keys or []
        mock.get_indexes.return_value = indexes or []
        return mock

    def test_create_table_includes_named_fk_constraint(self, tmp_path: Path) -> None:
        """CREATE TABLE incluye CONSTRAINT nombre FOREIGN KEY cuando hay FK con nombre."""
        cols = [
            {"name": "id",      "type": "INTEGER",     "nullable": False, "default": None},
            {"name": "user_id", "type": "INTEGER",      "nullable": True,  "default": None},
        ]
        fks = [
            {
                "name": "fk_orders_user",
                "constrained_columns": ["user_id"],
                "referred_table": "users",
                "referred_columns": ["id"],
            }
        ]
        inspector = self._make_full_inspector(cols, ["id"], foreign_keys=fks)
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="orders", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CONSTRAINT fk_orders_user FOREIGN KEY (user_id) REFERENCES users (id)" in content

    def test_create_table_includes_unnamed_fk_constraint(self, tmp_path: Path) -> None:
        """CREATE TABLE incluye FOREIGN KEY sin CONSTRAINT cuando la FK no tiene nombre."""
        cols = [
            {"name": "id",      "type": "INTEGER", "nullable": False, "default": None},
            {"name": "cat_id",  "type": "INTEGER", "nullable": True,  "default": None},
        ]
        fks = [
            {
                "name": "",
                "constrained_columns": ["cat_id"],
                "referred_table": "categories",
                "referred_columns": ["id"],
            }
        ]
        inspector = self._make_full_inspector(cols, ["id"], foreign_keys=fks)
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="products", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "FOREIGN KEY (cat_id) REFERENCES categories (id)" in content
        assert "CONSTRAINT  FOREIGN KEY" not in content

    def test_create_table_includes_secondary_index(self, tmp_path: Path) -> None:
        """CREATE TABLE genera CREATE INDEX para índices secundarios."""
        cols = [
            {"name": "id",    "type": "INTEGER",     "nullable": False, "default": None},
            {"name": "email", "type": "VARCHAR(200)", "nullable": True,  "default": None},
        ]
        indexes = [
            {"name": "idx_email", "column_names": ["email"], "unique": False}
        ]
        inspector = self._make_full_inspector(cols, ["id"], indexes=indexes)
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="users", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CREATE INDEX idx_email ON users (email)" in content

    def test_create_table_includes_unique_index(self, tmp_path: Path) -> None:
        """CREATE TABLE genera CREATE UNIQUE INDEX para índices únicos."""
        cols = [
            {"name": "id",    "type": "INTEGER",     "nullable": False, "default": None},
            {"name": "email", "type": "VARCHAR(200)", "nullable": False, "default": None},
        ]
        indexes = [
            {"name": "uq_email", "column_names": ["email"], "unique": True}
        ]
        inspector = self._make_full_inspector(cols, ["id"], indexes=indexes)
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="users", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CREATE UNIQUE INDEX uq_email ON users (email)" in content

    def test_create_table_index_after_closing_paren(self, tmp_path: Path) -> None:
        """El CREATE INDEX aparece después del cierre de CREATE TABLE (…)."""
        cols = [{"name": "id", "type": "INTEGER", "nullable": False, "default": None}]
        indexes = [{"name": "idx_id", "column_names": ["id"], "unique": False}]
        inspector = self._make_full_inspector(cols, ["id"], indexes=indexes)
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="t", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        close_paren_pos = content.index(");")
        index_pos = content.index("CREATE INDEX idx_id")
        assert index_pos > close_paren_pos

    def test_create_table_skips_fk_without_referred_table(self, tmp_path: Path) -> None:
        """FK sin tabla referenciada se omite del DDL."""
        cols = [{"name": "id", "type": "INTEGER", "nullable": False, "default": None}]
        fks = [
            {
                "name": "bad_fk",
                "constrained_columns": ["id"],
                "referred_table": "",   # vacío → debe omitirse
                "referred_columns": [],
            }
        ]
        inspector = self._make_full_inspector(cols, ["id"], foreign_keys=fks)
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="t", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=inspector).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "bad_fk" not in content
        assert "FOREIGN KEY" not in content

    def test_create_table_inspector_raises_on_fk_uses_empty(self, tmp_path: Path) -> None:
        """Si el Inspector lanza al obtener FKs, se omiten sin error."""
        cols = [{"name": "id", "type": "INTEGER", "nullable": False, "default": None}]
        mock = MagicMock()
        mock.get_columns.return_value = cols
        mock.get_pk_constraint.return_value = {"constrained_columns": ["id"], "name": None}
        mock.get_foreign_keys.side_effect = Exception("FK error")
        mock.get_indexes.return_value = []
        diff = SchemaDiff(
            source_label="A", target_label="B",
            table_diffs=(TableDiff(table_name="t", status="added"),),
        )
        output = tmp_path / "diff.sql"
        SqlDiffReporter(source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CREATE TABLE" in content
        assert "id INTEGER" in content


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter — nullable diff con tipo real (inspector)
# ---------------------------------------------------------------------------


class TestSqlDiffReporterNullableWithInspector:
    """Verifica que _render_nullable_diff usa el tipo real del Inspector."""

    @staticmethod
    def _nullable_diff(table: str, col: str, source_nullable: str) -> SchemaDiff:
        return SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name=table,
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name=col,
                            attribute="nullable",
                            source_value=source_nullable,
                            target_value="False",
                        ),
                    ),
                ),
            ),
        )

    def test_mariadb_nullable_uses_real_type(self, tmp_path: Path) -> None:
        """MariaDB: MODIFY COLUMN usa el tipo real del Inspector."""
        mock = MagicMock()
        mock.get_columns.return_value = [
            {"name": "email", "type": "VARCHAR(150)", "nullable": True, "default": None}
        ]
        diff = self._nullable_diff("users", "email", "True")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb", source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN email VARCHAR(150) NULL" in content
        assert "<TIPO>" not in content

    def test_mariadb_not_null_uses_real_type(self, tmp_path: Path) -> None:
        """MariaDB: MODIFY COLUMN usa NOT NULL con tipo real."""
        mock = MagicMock()
        mock.get_columns.return_value = [
            {"name": "code", "type": "CHAR(10)", "nullable": False, "default": None}
        ]
        diff = self._nullable_diff("items", "code", "False")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb", source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN code CHAR(10) NOT NULL" in content

    def test_mariadb_nullable_fallback_placeholder_without_inspector(self, tmp_path: Path) -> None:
        """MariaDB sin inspector → <TIPO> placeholder."""
        diff = self._nullable_diff("users", "email", "True")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN email <TIPO>" in content

    def test_oracle_nullable_uses_real_type(self, tmp_path: Path) -> None:
        """Oracle: MODIFY usa tipo real cuando inspector disponible."""
        mock = MagicMock()
        mock.get_columns.return_value = [
            {"name": "status", "type": "VARCHAR2(20)", "nullable": True, "default": None}
        ]
        diff = self._nullable_diff("orders", "status", "False")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle", source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY (status VARCHAR2(20) NOT NULL)" in content

    def test_oracle_nullable_fallback_without_type(self, tmp_path: Path) -> None:
        """Oracle sin inspector → MODIFY sin tipo (solo NULL/NOT NULL)."""
        diff = self._nullable_diff("orders", "status", "False")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY (status NOT NULL)" in content

    def test_postgresql_nullable_uses_alter_column(self, tmp_path: Path) -> None:
        """PostgreSQL usa ALTER COLUMN DROP/SET NOT NULL (no MODIFY COLUMN)."""
        diff = self._nullable_diff("users", "email", "False")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER COLUMN email SET NOT NULL" in content
        assert "MODIFY COLUMN" not in content


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter — default diff dialect-aware
# ---------------------------------------------------------------------------


class TestSqlDiffReporterDefaultDialectAware:
    """Verifica que _render_default_diff genera sintaxis dialect-específica."""

    @staticmethod
    def _default_diff(table: str, col: str, new_default: str) -> SchemaDiff:
        return SchemaDiff(
            source_label="A",
            target_label="B",
            table_diffs=(
                TableDiff(
                    table_name=table,
                    status="modified",
                    column_diffs=(
                        ColumnDiff(
                            column_name=col,
                            attribute="default",
                            source_value=new_default,
                            target_value="",
                        ),
                    ),
                ),
            ),
        )

    def test_mariadb_set_default_uses_modify_column(self, tmp_path: Path) -> None:
        """MariaDB: SET DEFAULT → MODIFY COLUMN col TIPO DEFAULT val."""
        mock = MagicMock()
        mock.get_columns.return_value = [
            {"name": "status", "type": "VARCHAR(20)", "nullable": True, "default": None}
        ]
        diff = self._default_diff("orders", "status", "'active'")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb", source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN status VARCHAR(20) DEFAULT 'active'" in content

    def test_mariadb_drop_default_uses_modify_column(self, tmp_path: Path) -> None:
        """MariaDB: DROP DEFAULT → MODIFY COLUMN sin DEFAULT."""
        mock = MagicMock()
        mock.get_columns.return_value = [
            {"name": "flag", "type": "TINYINT", "nullable": True, "default": None}
        ]
        diff = self._default_diff("t", "flag", "None")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb", source_inspector=mock).render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY COLUMN flag TINYINT" in content
        assert "DEFAULT eliminado" in content

    def test_oracle_set_default_uses_modify(self, tmp_path: Path) -> None:
        """Oracle: SET DEFAULT → MODIFY (col DEFAULT val)."""
        diff = self._default_diff("products", "active", "1")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY (active DEFAULT 1)" in content

    def test_oracle_drop_default_uses_modify_null(self, tmp_path: Path) -> None:
        """Oracle: DROP DEFAULT → MODIFY (col DEFAULT NULL)."""
        diff = self._default_diff("products", "active", "NULL")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "MODIFY (active DEFAULT NULL)" in content

    def test_postgresql_set_default_uses_set_default(self, tmp_path: Path) -> None:
        """PostgreSQL: SET DEFAULT → ALTER COLUMN col SET DEFAULT val."""
        diff = self._default_diff("users", "score", "0")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER COLUMN score SET DEFAULT 0" in content

    def test_postgresql_drop_default_uses_drop_default(self, tmp_path: Path) -> None:
        """PostgreSQL: DROP DEFAULT → ALTER COLUMN col DROP DEFAULT."""
        diff = self._default_diff("users", "score", "")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER COLUMN score DROP DEFAULT" in content

    def test_generic_set_default(self, tmp_path: Path) -> None:
        """Generic: SET DEFAULT usa ALTER COLUMN SET DEFAULT."""
        diff = self._default_diff("t", "col", "42")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="generic").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "ALTER COLUMN col SET DEFAULT 42" in content


# ---------------------------------------------------------------------------
# Tests de DdlNormalizer
# ---------------------------------------------------------------------------


class TestDdlNormalizerViews:
    """Unit tests para DdlNormalizer.normalize_view_ddl y extract_view_body."""

    _MARIADB_FULL = (
        "CREATE ALGORITHM=UNDEFINED DEFINER=`root`@`localhost` "
        "SQL SECURITY DEFINER VIEW `my_view` AS SELECT id FROM t"
    )
    _MARIADB_RESULT = "CREATE VIEW `my_view` AS SELECT id FROM t"

    def test_strips_algorithm(self) -> None:
        raw = "CREATE ALGORITHM=TEMPTABLE DEFINER=`u`@`h` VIEW v AS SELECT 1"
        result = DdlNormalizer.normalize_view_ddl(raw)
        assert "ALGORITHM" not in result

    def test_strips_definer_backticks(self) -> None:
        result = DdlNormalizer.normalize_view_ddl(self._MARIADB_FULL)
        assert "DEFINER" not in result

    def test_strips_sql_security(self) -> None:
        result = DdlNormalizer.normalize_view_ddl(self._MARIADB_FULL)
        assert "SQL SECURITY" not in result

    def test_full_mariadb_normalization(self) -> None:
        """Todos los artefactos eliminados; queda CREATE VIEW `name` AS ..."""
        result = DdlNormalizer.normalize_view_ddl(self._MARIADB_FULL)
        assert result == self._MARIADB_RESULT

    def test_strips_definer_single_quotes(self) -> None:
        raw = "CREATE DEFINER='admin'@'%' VIEW v AS SELECT 1"
        result = DdlNormalizer.normalize_view_ddl(raw)
        assert "DEFINER" not in result
        assert "SELECT 1" in result

    def test_strips_definer_unquoted(self) -> None:
        raw = "CREATE DEFINER=admin@localhost VIEW v AS SELECT 1"
        result = DdlNormalizer.normalize_view_ddl(raw)
        assert "DEFINER" not in result

    def test_strips_oracle_force(self) -> None:
        raw = "CREATE OR REPLACE FORCE VIEW my_view AS SELECT id FROM t"
        result = DdlNormalizer.normalize_view_ddl(raw)
        assert "FORCE" not in result
        assert "CREATE OR REPLACE VIEW my_view" in result

    def test_noop_on_clean_ddl(self) -> None:
        clean = "CREATE OR REPLACE VIEW v AS SELECT id FROM t"
        result = DdlNormalizer.normalize_view_ddl(clean)
        assert result == clean

    def test_case_insensitive(self) -> None:
        raw = "create algorithm=undefined definer=`u`@`h` sql security definer view v as select 1"
        result = DdlNormalizer.normalize_view_ddl(raw)
        assert "algorithm" not in result.lower() or "algorithm" not in result
        assert "definer" not in result.lower()

    def test_extract_body_from_full_ddl(self) -> None:
        body = DdlNormalizer.extract_view_body(self._MARIADB_FULL)
        assert body == "SELECT id FROM t"

    def test_extract_body_from_backtick_name(self) -> None:
        raw = "CREATE OR REPLACE VIEW `my view` AS SELECT a, b FROM t WHERE x = 1"
        body = DdlNormalizer.extract_view_body(raw)
        assert body == "SELECT a, b FROM t WHERE x = 1"

    def test_extract_body_returns_input_when_no_header(self) -> None:
        select_body = "SELECT id FROM t"
        result = DdlNormalizer.extract_view_body(select_body)
        assert result == select_body

    def test_extract_body_double_quoted_name(self) -> None:
        raw = 'CREATE VIEW "my_view" AS SELECT 1'
        body = DdlNormalizer.extract_view_body(raw)
        assert body == "SELECT 1"


class TestDdlNormalizerRoutines:
    """Unit tests para DdlNormalizer.normalize_routine_ddl."""

    def test_strips_definer_from_procedure(self) -> None:
        raw = "CREATE DEFINER=`root`@`localhost` PROCEDURE my_proc() BEGIN SELECT 1; END"
        result = DdlNormalizer.normalize_routine_ddl(raw)
        assert "DEFINER" not in result
        assert "PROCEDURE my_proc" in result

    def test_preserves_language_clause(self) -> None:
        raw = "CREATE FUNCTION my_fn() RETURNS void LANGUAGE plpgsql AS $$ BEGIN END $$"
        result = DdlNormalizer.normalize_routine_ddl(raw)
        assert "LANGUAGE plpgsql" in result

    def test_noop_on_clean_routine(self) -> None:
        clean = "CREATE PROCEDURE p() BEGIN END"
        result = DdlNormalizer.normalize_routine_ddl(clean)
        assert result == clean

    def test_strips_definer_single_quotes(self) -> None:
        raw = "CREATE DEFINER='u'@'h' FUNCTION f() RETURNS INT RETURN 1"
        result = DdlNormalizer.normalize_routine_ddl(raw)
        assert "DEFINER" not in result


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter — normalización de vistas
# ---------------------------------------------------------------------------


class TestSqlDiffReporterViewNormalization:
    """Tests de _render_single_view_diff con DDL con artefactos MariaDB."""

    _MARIADB_VIEW_DDL = (
        "CREATE ALGORITHM=UNDEFINED DEFINER=`root`@`localhost` "
        "SQL SECURITY DEFINER VIEW `orders_view` AS SELECT id, total FROM orders"
    )

    @staticmethod
    def _view_diff(status: str, source_def: str = "", target_def: str = "") -> SchemaDiff:
        return SchemaDiff(
            source_label="src",
            target_label="tgt",
            view_diffs=(
                ViewDiff(
                    view_name="orders_view",
                    status=status,  # type: ignore[arg-type]
                    source_definition=source_def,
                    target_definition=target_def,
                ),
            ),
        )

    def test_added_view_no_double_wrap(self, tmp_path: Path) -> None:
        """La DDL no debe contener 'CREATE VIEW ... AS\\nCREATE VIEW ...'."""
        diff = self._view_diff("added", source_def=self._MARIADB_VIEW_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        # El cuerpo CREATE VIEW no debe aparecer duplicado
        assert content.count("CREATE") == content.count("VIEW") or "CREATE VIEW orders_view AS\nCREATE" not in content

    def test_added_view_definer_stripped(self, tmp_path: Path) -> None:
        diff = self._view_diff("added", source_def=self._MARIADB_VIEW_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DEFINER" not in content
        assert "ALGORITHM" not in content
        assert "SQL SECURITY" not in content

    def test_added_view_body_present(self, tmp_path: Path) -> None:
        diff = self._view_diff("added", source_def=self._MARIADB_VIEW_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="generic").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "SELECT id, total FROM orders" in content

    def test_modified_view_uses_create_or_replace(self, tmp_path: Path) -> None:
        diff = self._view_diff("modified", source_def=self._MARIADB_VIEW_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "CREATE OR REPLACE VIEW" in content
        assert "DEFINER" not in content

    def test_modified_view_sqlserver_uses_drop_go_create(self, tmp_path: Path) -> None:
        diff = self._view_diff("modified", source_def="CREATE VIEW orders_view AS SELECT id FROM orders")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="sqlserver").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DROP VIEW" in content
        assert "GO" in content
        assert "CREATE VIEW orders_view" in content

    def test_modified_view_oracle_force_stripped(self, tmp_path: Path) -> None:
        oracle_ddl = "CREATE OR REPLACE FORCE VIEW orders_view AS SELECT id FROM orders"
        diff = self._view_diff("modified", source_def=oracle_ddl)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="oracle").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "FORCE" not in content
        assert "CREATE OR REPLACE VIEW orders_view" in content

    def test_removed_view_emits_drop(self, tmp_path: Path) -> None:
        diff = self._view_diff("removed")
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="postgresql").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DROP VIEW" in content
        assert "orders_view" in content


# ---------------------------------------------------------------------------
# Tests de SqlDiffReporter — normalización de rutinas
# ---------------------------------------------------------------------------


class TestSqlDiffReporterRoutineNormalization:
    """Tests de _render_single_routine_diff con DDL con DEFINER."""

    _MARIADB_PROC_DDL = (
        "CREATE DEFINER=`root`@`localhost` PROCEDURE calc_total()\n"
        "BEGIN\n  SELECT SUM(amount) FROM orders;\nEND"
    )

    @staticmethod
    def _routine_diff(status: str, source_def: str = "", target_def: str = "") -> SchemaDiff:
        return SchemaDiff(
            source_label="src",
            target_label="tgt",
            routine_diffs=(
                RoutineDiff(
                    routine_name="calc_total",
                    routine_type="PROCEDURE",
                    status=status,  # type: ignore[arg-type]
                    source_definition=source_def,
                    target_definition=target_def,
                ),
            ),
        )

    def test_added_routine_definer_stripped(self, tmp_path: Path) -> None:
        diff = self._routine_diff("added", source_def=self._MARIADB_PROC_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DEFINER" not in content
        assert "PROCEDURE calc_total" in content

    def test_modified_routine_definer_stripped(self, tmp_path: Path) -> None:
        diff = self._routine_diff("modified", source_def=self._MARIADB_PROC_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="mariadb").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "DEFINER" not in content
        assert "SELECT SUM(amount) FROM orders" in content

    def test_routine_body_preserved(self, tmp_path: Path) -> None:
        diff = self._routine_diff("added", source_def=self._MARIADB_PROC_DDL)
        output = tmp_path / "diff.sql"
        SqlDiffReporter(dialect="generic").render(diff, output)
        content = output.read_text(encoding="utf-8")
        assert "SELECT SUM(amount) FROM orders" in content
