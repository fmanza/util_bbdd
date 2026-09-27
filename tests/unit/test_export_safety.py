"""Escritura segura de informes (``_core.export_safety``): caracteres de control y fórmulas en Excel y Word."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

import docx
import openpyxl
import pandas as pd
import pytest

from openpyxl.worksheet.worksheet import Worksheet

from util_bbdd._core.export_safety import SafeWorkbook, SafeWorksheet, safe_frame, safe_text, safe_value
from util_bbdd.comparar.models import SchemaDiff
from util_bbdd.comparar.reporters.excel_reporter import ExcelDiffReporter
from util_bbdd.comparar.reporters.word_reporter import WordDiffReporter
from util_bbdd.doc.infrastructure.renderers.excel_generator import ExcelGenerator
from util_bbdd.doc.infrastructure.renderers.word_generator import WordGenerator

pytestmark = pytest.mark.unit

_FORMULA = '=HYPERLINK("http://example.invalid","clic")'


def _all_cells(path: Path) -> list[openpyxl.cell.cell.Cell]:
    workbook = openpyxl.load_workbook(path)
    return [cell for ws in workbook.worksheets for row in ws.iter_rows() for cell in row]


class TestSafeValue:
    def test_removes_control_characters_and_keeps_tab_and_newline(self) -> None:
        assert safe_value("a\x00b\x0bc\x1bd\te\nf") == "abcd\te\nf"

    @pytest.mark.parametrize("value", [None, 3, 2.5, True])
    def test_non_strings_are_returned_unchanged(self, value: object) -> None:
        assert safe_value(value) is value

    @pytest.mark.parametrize(("value", "expected"), [("a\x0bb", "ab"), (None, ""), (12, "12")])
    def test_safe_text_returns_xml_compatible_text(self, value: object, expected: str) -> None:
        assert safe_text(value) == expected

    def test_safe_frame_cleans_text_columns_only_and_copies(self) -> None:
        df = pd.DataFrame({"tabla": ["x\x07y"], "filas": [3]})

        result = safe_frame(df)

        assert result.loc[0, "tabla"] == "xy"
        assert result.loc[0, "filas"] == 3
        assert df.loc[0, "tabla"] == "x\x07y"


class TestSafeWorkbook:
    def test_every_sheet_is_a_safe_worksheet(self) -> None:
        wb = SafeWorkbook()
        wb.create_sheet("Otra")

        assert all(type(ws) is SafeWorksheet for ws in wb.worksheets)
        assert all(isinstance(ws, Worksheet) for ws in wb.worksheets)

    def test_cell_append_and_setitem_remove_control_characters(self, tmp_path: Path) -> None:
        wb = SafeWorkbook()
        ws = wb.active
        assert ws is not None
        ws.cell(row=1, column=1, value="uno\x00")
        ws.append(["dos\x0b", 2])
        ws.append({"C": "tres\x1b"})
        ws["D1"] = "cuatro\x0c"
        wb.save(tmp_path / "r.xlsx")

        values = [cell.value for cell in _all_cells(tmp_path / "r.xlsx") if cell.value is not None]
        assert values == ["uno", "cuatro", "dos", 2, "tres"]

    def test_write_only_is_rejected(self) -> None:
        # Sus hojas (WriteOnlyWorksheet) no pasan por cell/append de Worksheet: no se podrían sanear.
        with pytest.raises(ValueError, match="write_only"):
            SafeWorkbook(write_only=True)

    def test_cell_without_value_keeps_the_existing_value(self) -> None:
        ws = SafeWorkbook().active
        assert ws is not None
        ws.cell(row=1, column=1, value="dato")

        assert ws.cell(row=1, column=1).value == "dato"

    def test_formulas_are_saved_as_text(self, tmp_path: Path) -> None:
        wb = SafeWorkbook()
        wb.create_sheet("Otra").cell(row=1, column=1, value=_FORMULA)
        wb.save(tmp_path / "r.xlsx")

        cells = _all_cells(tmp_path / "r.xlsx")
        assert [(c.data_type, c.value) for c in cells if c.value] == [("s", _FORMULA)]


class TestExcelDiffReporter:
    def test_catalog_values_with_control_characters_or_formulas_are_written_safely(self, tmp_path: Path) -> None:
        # Antes: IllegalCharacterError con "\x0b" y la etiqueta "=HYPERLINK(...)" quedaba como fórmula ejecutable.
        diff = SchemaDiff(source_label=_FORMULA, target_label="destino\x0b")

        path = ExcelDiffReporter().render(diff, tmp_path / "diff.xlsx")

        cells = _all_cells(path)
        assert not [c.coordinate for c in cells if c.data_type == "f"]
        values = [c.value for c in cells]
        assert _FORMULA in values
        assert "destino" in values


class TestWordDiffReporter:
    def test_catalog_values_with_control_characters_do_not_abort_the_report(self, tmp_path: Path) -> None:
        # python-docx: "ValueError: All strings must be XML compatible" con "\x0b" o "\x00".
        diff = SchemaDiff(source_label="origen\x0b", target_label="destino\x00", summary="x\x1by")

        path = WordDiffReporter().render(diff, tmp_path / "diff.docx")

        cells = [c.text for t in docx.Document(str(path)).tables for r in t.rows for c in r.cells]
        assert "origen" in cells
        assert "destino" in cells


class TestDocGenerators:
    """P5: comentarios y nombres del catálogo con caracteres de control o fórmulas en ``doc/`` (Excel y Word)."""

    def test_excel_table_header_is_written_safely(self, tmp_path: Path) -> None:
        wb = SafeWorkbook()
        ws = wb.active
        assert ws is not None
        generator = ExcelGenerator(database="db", headers={})

        generator._process_table(
            ws, 1, "=tabla\x0b", "comentario\x00", "utf8\x1b", cast(Any, object()), [], [False, False]
        )
        wb.save(tmp_path / "doc.xlsx")

        values = [c.value for c in _all_cells(tmp_path / "doc.xlsx") if c.value]
        assert values == ["Tabla: =tabla", "Comentario: comentario", "Charset & Collation: utf8"]

    def test_excel_raw_catalog_value_starting_with_equals_is_text(self, tmp_path: Path) -> None:
        wb = SafeWorkbook()
        ws = wb.active
        assert ws is not None
        ws.append(["col", _FORMULA])  # como las filas de columnas del catálogo (default o comentario)
        wb.save(tmp_path / "doc.xlsx")

        assert [(c.data_type, c.value) for c in _all_cells(tmp_path / "doc.xlsx") if c.value == _FORMULA] == [
            ("s", _FORMULA)
        ]

    def test_word_table_section_is_written_safely(self, tmp_path: Path) -> None:
        document = docx.Document()
        generator = WordGenerator(database="db", headers={})

        generator._add_table_section(
            document, "tabla\x0b", "comentario\x00", cast(Any, object()), [False] * 8, charset_collation="utf8\x1b"
        )
        document.save(str(tmp_path / "doc.docx"))

        text = " | ".join(p.text for p in docx.Document(str(tmp_path / "doc.docx")).paragraphs)
        assert "Tabla: tabla" in text
        assert "Comentario: comentario" in text
        assert "Charset & Collation: utf8" in text
