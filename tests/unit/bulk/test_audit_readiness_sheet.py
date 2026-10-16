"""Hoja «Condiciones_motor» del Excel de auditoría de bulk (AuditManager.exportar_resumen_excel)."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest

from sqlalchemy import create_engine, event, text
from sqlalchemy.engine import Engine

from util_bbdd import BulkCrud
from util_bbdd.bulk import AuditManager
from util_bbdd.bulk.options import AuditOptions
from util_bbdd.connection.performance_advisor import EngineReadinessAdvisor

SHEET = "Condiciones_motor"


@pytest.fixture
def engine(tmp_path: Path) -> Iterator[Engine]:
    eng = create_engine(f"sqlite:///{tmp_path / 'db.sqlite'}")
    with eng.begin() as conn:
        conn.execute(text("CREATE TABLE t (id INTEGER PRIMARY KEY, n TEXT)"))
    yield eng
    eng.dispose()


def _load(engine: Engine, audit_csv: Path) -> BulkCrud:
    crud = BulkCrud(engine, audit=AuditOptions(csv_path=audit_csv))
    crud.bulk_insert_from_df(pd.DataFrame({"id": [1, 2], "n": ["a", "b"]}), "t", generar_sql=False)
    return crud


def _first_column(path: Path) -> list[object]:
    return [row[0].value for row in openpyxl.load_workbook(path)[SHEET].iter_rows()]


def test_bulkcrud_export_ends_with_the_engine_conditions_sheet(engine: Engine, tmp_path: Path) -> None:
    crud = _load(engine, tmp_path / "audit.csv")
    checkouts: list[object] = []
    event.listen(engine, "checkout", lambda *args: checkouts.append(args))

    path = crud.audit_manager.exportar_resumen_excel()

    assert openpyxl.load_workbook(path).sheetnames[-2:] == ["Resumen_global", SHEET]
    column = _first_column(path)
    assert f"Engine · {EngineReadinessAdvisor.describe_engine(engine)}" in column
    assert "journal_mode" in [row[1].value for row in openpyxl.load_workbook(path)[SHEET].iter_rows()]
    assert checkouts == []  # la verificación de BulkCrud.__init__, sin nueva conexión


def test_standalone_audit_manager_adds_the_sheet_only_with_an_engine(engine: Engine, tmp_path: Path) -> None:
    _load(engine, tmp_path / "audit.csv")
    manager = AuditManager(tmp_path / "audit.csv")

    assert SHEET not in openpyxl.load_workbook(manager.exportar_resumen_excel()).sheetnames
    assert openpyxl.load_workbook(manager.exportar_resumen_excel(engine=engine)).sheetnames[-1] == SHEET


def test_a_failed_verification_does_not_stop_the_export(engine: Engine, tmp_path: Path) -> None:
    crud = _load(engine, tmp_path / "audit.csv")
    with patch.object(EngineReadinessAdvisor, "engine_report", return_value=None):
        path = crud.audit_manager.exportar_resumen_excel()

    assert "⚠️ No se pudieron verificar las condiciones de este motor." in _first_column(path)
