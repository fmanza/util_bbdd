"""Exportación de auditoría a Excel: distingue generación de SQL (``sql=Sí``) de ejecución en BD (``sql=No``).

``tiempo_seg`` es la única columna de tiempo y mide cosas distintas según ``sql``; si el resumen las agrupara juntas,
``tiempo_total`` y ``tiempo_medio`` mezclarían latencia de BD con tiempo de generación de SQL.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from util_bbdd.bulk.audit import AuditManager

pytestmark = pytest.mark.unit

_GLOBAL_TABLA_FECHA = ["fecha", "tabla", "operacion", "sql", "total_filas", "tiempo_total", "lotes", "tiempo_medio"]


@pytest.fixture
def manager(tmp_path: Path) -> AuditManager:
    mgr = AuditManager(audit_csv_path=tmp_path / "audit.csv")
    mgr.registrar_operacion("orders", "sqlite", "INSERT", 10, 1.5, False)
    mgr.registrar_operacion("orders", "sqlite", "INSERT", 10, 0.4, True)
    mgr.registrar_operacion("orders", "sqlite", "INSERT", 10, 0.6, True)
    mgr.registrar_operacion("orders", "sqlite", "UPDATE", 5, 2.0, False)
    return mgr


def _sheets(path: Path) -> dict[str, pd.DataFrame]:
    return pd.read_excel(path, sheet_name=None)


@pytest.mark.parametrize("agrupar_por", ["tabla", "fecha", "tabla_fecha"])
def test_every_sheet_has_sql_column(manager: AuditManager, agrupar_por: str) -> None:
    sheets = _sheets(manager.exportar_resumen_excel(agrupar_por))  # type: ignore[arg-type]

    assert set(sheets) == {f"Insert_{agrupar_por}", f"Update_{agrupar_por}", "Resumen_global"}
    assert all("sql" in df.columns for df in sheets.values())


def test_generation_and_execution_are_separate_rows_with_their_own_aggregates(manager: AuditManager) -> None:
    insert = _sheets(manager.exportar_resumen_excel("tabla"))["Insert_tabla"].set_index("sql")

    assert set(insert.index) == {"No", "Sí"}
    assert insert.loc["No", ["total_filas", "tiempo_total", "lotes", "tiempo_medio"]].tolist() == [10, 1.5, 1, 1.5]
    assert insert.loc["Sí", ["total_filas", "tiempo_total", "lotes", "tiempo_medio"]].tolist() == [20, 1.0, 2, 0.5]


def test_resumen_global_keeps_sql_column_in_tabla_fecha_mode(manager: AuditManager) -> None:
    global_sheet = _sheets(manager.exportar_resumen_excel("tabla_fecha"))["Resumen_global"]

    assert list(global_sheet.columns) == _GLOBAL_TABLA_FECHA
    assert sorted(global_sheet["sql"].tolist()) == ["No", "No", "Sí"]


def test_row_without_sql_value_is_not_dropped_from_the_summary(manager: AuditManager, tmp_path: Path) -> None:
    with open(tmp_path / "audit.csv", "a", newline="", encoding="utf-8") as handle:
        handle.write("2026-09-25 10:00:00\torders\tsqlite\tINSERT\t7\t0,100\t\n")

    insert = _sheets(manager.exportar_resumen_excel("tabla"))["Insert_tabla"]

    assert insert["total_filas"].sum() == 10 + 20 + 7


def test_text_starting_with_equals_is_written_as_text_not_formula(tmp_path: Path) -> None:
    # H12 (audit-code 2026-09-25): openpyxl guarda como fórmula toda cadena que empieza por "=".
    from openpyxl import load_workbook

    formula = '=HYPERLINK("http://example.invalid","clic")'
    mgr = AuditManager(audit_csv_path=tmp_path / "audit.csv")
    mgr.registrar_operacion(formula, "sqlite", "INSERT", 1, 0.1, False)

    workbook = load_workbook(mgr.exportar_resumen_excel("tabla"))

    cells = [cell for ws in workbook.worksheets for row in ws.iter_rows() for cell in row]
    assert not [cell.coordinate for cell in cells if cell.data_type == "f"]
    assert sum(cell.value == formula for cell in cells) == len(workbook.worksheets)


def test_control_characters_in_a_table_name_do_not_abort_the_export(tmp_path: Path) -> None:
    # openpyxl lanza IllegalCharacterError con caracteres de control; se eliminan (safe_frame) antes de to_excel.
    from openpyxl import load_workbook

    mgr = AuditManager(audit_csv_path=tmp_path / "audit.csv")
    mgr.registrar_operacion("tabla\x1brara", "sqlite", "INSERT", 1, 0.1, False)

    workbook = load_workbook(mgr.exportar_resumen_excel("tabla"))

    assert "tablarara" in [cell.value for row in workbook["Resumen_global"].iter_rows() for cell in row]
