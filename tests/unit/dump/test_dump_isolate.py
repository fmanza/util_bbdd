"""``DumpConfig(on_error="isolate", rejects_csv_path, max_rejects)`` (Fase 4, sesión 2026-10-01u).

Las filas que el destino no admite (aquí, NOT NULL y PK duplicada) se apartan al CSV de rechazos y el resto de la tabla
se copia; la tabla queda ``success`` con ``rows_rejected``.
"""

from __future__ import annotations

import csv

from pathlib import Path

import openpyxl
import pytest

from sqlalchemy import Engine, create_engine, text

from util_bbdd.dump import DumpConfig, DumpOrchestrator

pytestmark = pytest.mark.unit


def _engines(tmp_path: Path) -> tuple[Engine, Engine]:
    src = create_engine(f"sqlite:///{tmp_path / 'src.db'}")
    dst = create_engine(f"sqlite:///{tmp_path / 'dst.db'}")
    with src.begin() as conn:
        conn.execute(text("CREATE TABLE a (id INTEGER PRIMARY KEY, v VARCHAR(10))"))
        conn.execute(text("CREATE TABLE b (id INTEGER PRIMARY KEY, v VARCHAR(10))"))
        conn.execute(text("INSERT INTO a VALUES (1, 'x'), (2, NULL), (3, 'y'), (4, NULL), (5, 'z')"))
        conn.execute(text("INSERT INTO b VALUES (1, 'p'), (2, NULL), (3, 'q')"))
    with dst.begin() as conn:  # el destino no admite nulos en v
        conn.execute(text("CREATE TABLE a (id INTEGER PRIMARY KEY, v VARCHAR(10) NOT NULL)"))
        conn.execute(text("CREATE TABLE b (id INTEGER PRIMARY KEY, v VARCHAR(10) NOT NULL)"))
    return src, dst


def _ids(engine: Engine, table: str) -> list[int]:
    with engine.connect() as conn:
        return [r[0] for r in conn.execute(text(f"SELECT id FROM {table} ORDER BY id"))]


def _rejects(path: Path) -> list[tuple[str, str, int]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return [(r["tabla"], r["particion"], int(r["fila"])) for r in csv.DictReader(handle, delimiter="\t")]


def _config(src: Engine, dst: Engine, tmp_path: Path, **kwargs: object) -> DumpConfig:
    return DumpConfig(
        source_engine=src,
        target_engine=dst,
        output_dir=tmp_path,
        on_error="isolate",
        rejects_csv_path=tmp_path / "rejects.csv",
        **kwargs,  # type: ignore[arg-type]
    )


def test_rejected_rows_are_set_aside_and_the_rest_is_copied(tmp_path: Path) -> None:
    src, dst = _engines(tmp_path)

    result = DumpOrchestrator(_config(src, dst, tmp_path, tables=["a", "b"])).run()

    by_table = {r.table_name: r for r in result.table_results}
    assert {name: (r.status, r.rows_rejected) for name, r in by_table.items()} == {
        "a": ("success", 2),
        "b": ("success", 1),
    }
    assert result.rows_rejected == 3
    assert _ids(dst, "a") == [1, 3, 5] and _ids(dst, "b") == [1, 3]
    assert sorted(_rejects(tmp_path / "rejects.csv")) == [("a", "", 1), ("a", "", 3), ("b", "", 1)]
    ws = openpyxl.load_workbook(result.report)["Auditoría Tablas"]
    assert ws.cell(1, 7).value == "Filas Rechazadas"
    assert {ws.cell(row, 1).value: ws.cell(row, 7).value for row in (2, 3)} == {"a": 2, "b": 1}
    src.dispose()
    dst.dispose()


def test_duplicate_keys_without_truncate(tmp_path: Path) -> None:
    src, dst = _engines(tmp_path)
    with src.begin() as conn:
        conn.execute(text("UPDATE a SET v = 'n' WHERE v IS NULL"))
    with dst.begin() as conn:
        conn.execute(text("INSERT INTO a VALUES (3, 'ya'), (5, 'ya')"))

    result = DumpOrchestrator(_config(src, dst, tmp_path, tables=["a"], truncate_target=False)).run()

    assert [(r.status, r.rows_rejected) for r in result.table_results] == [("success", 2)]
    assert _ids(dst, "a") == [1, 2, 3, 4, 5]
    src.dispose()
    dst.dispose()


def test_max_rejects_leaves_the_table_in_error_and_the_dump_goes_on(tmp_path: Path) -> None:
    src, dst = _engines(tmp_path)

    result = DumpOrchestrator(_config(src, dst, tmp_path, tables=["a", "b"], max_rejects=1)).run()

    by_table = {r.table_name: r for r in result.table_results}
    assert by_table["a"].status == "error"
    assert "NOT NULL" in by_table["a"].error_msg
    assert by_table["a"].rows_rejected == 1  # B-1: el rechazo escrito cuenta aunque no se copiara nada de la tabla
    assert (by_table["b"].status, by_table["b"].rows_rejected) == ("success", 1)
    assert _ids(dst, "a") == [] and _ids(dst, "b") == [1, 3]
    src.dispose()
    dst.dispose()


def test_raise_is_unchanged(tmp_path: Path) -> None:
    src, dst = _engines(tmp_path)

    result = DumpOrchestrator(DumpConfig(source_engine=src, target_engine=dst, output_dir=tmp_path, tables=["a"])).run()

    assert [(r.status, r.rows_rejected) for r in result.table_results] == [("error", 0)]
    assert result.rows_rejected == 0
    src.dispose()
    dst.dispose()


@pytest.mark.parametrize(
    "kwargs",
    [
        {"on_error": "skip"},
        {"on_error": "isolate"},
        {"rejects_csv_path": Path("r.csv")},
        {"on_error": "isolate", "rejects_csv_path": Path("r.csv"), "max_rejects": -1},
    ],
)
def test_invalid_parameters_fail_before_touching_the_target(tmp_path: Path, kwargs: dict[str, object]) -> None:
    src, dst = _engines(tmp_path)
    with pytest.raises(ValueError):
        DumpConfig(source_engine=src, target_engine=dst, **kwargs)  # type: ignore[arg-type]
    src.dispose()
    dst.dispose()
