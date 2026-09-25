"""Tests unitarios para DumpReporter (generación Excel 3 hojas)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from openpyxl.worksheet.worksheet import Worksheet

from util_bbdd._core.exceptions import DumpReportError
from util_bbdd.dump._constraints import ConstraintResult
from util_bbdd.dump._models import DumpAuditSummary, PKSyncResult, TableDumpResult
from util_bbdd.dump.reporters import DumpReporter

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _table_result(
    name: str = "orders",
    status: str = "success",
    source_rows: int = 100,
    target_after: int = 100,
    skipped: list[str] | None = None,
    in_cycle: bool = False,
    error_msg: str = "",
) -> TableDumpResult:
    return TableDumpResult(
        table_name=name,
        source_rows=source_rows,
        target_rows_before=0,
        target_rows_after=target_after,
        status=status,  # type: ignore[arg-type]
        skipped_columns=skipped or [],
        load_order_position=0,
        in_cycle=in_cycle,
        error_msg=error_msg,
    )


def _constraint_result(status: str = "enabled") -> ConstraintResult:
    return ConstraintResult(
        table_name="orders",
        constraint_name="fk_orders_cust",
        dialect="mssql",
        status=status,  # type: ignore[arg-type]
        error_msg="some error" if status == "failed" else "",
    )


# ---------------------------------------------------------------------------
# DumpReporter.generate — estructura del fichero
# ---------------------------------------------------------------------------


class TestDumpReporterSheetStructure:
    def test_generates_xlsx_with_three_sheets(self, tmp_path: Path):
        reporter = DumpReporter(tmp_path)
        path = reporter.generate(
            table_results=[_table_result()],
            missing_tables=[],
            constraint_results=[],
            cycles=[],
        )

        assert path.exists()
        wb = openpyxl.load_workbook(path)
        assert "Auditoría Tablas" in wb.sheetnames
        assert "Tablas Faltantes" in wb.sheetnames
        assert "Columnas Inexistentes" in wb.sheetnames

    def test_output_dir_created_if_missing(self, tmp_path: Path):
        new_dir = tmp_path / "subdir" / "reports"
        reporter = DumpReporter(new_dir)
        path = reporter.generate(
            table_results=[],
            missing_tables=[],
            constraint_results=[],
            cycles=[],
        )
        assert path.exists()

    def test_filename_contains_timestamp(self, tmp_path: Path):
        reporter = DumpReporter(tmp_path)
        path = reporter.generate(
            table_results=[],
            missing_tables=[],
            constraint_results=[],
            cycles=[],
        )
        assert "dump_report_" in path.name
        assert path.suffix == ".xlsx"


class TestDumpReporterDumpReportError:
    """FND-03/FND-04 (AUDIT_CODE 2026-09-22, decisión confirmada por el usuario):
    un fallo de I/O generando el informe (disco lleno, permisos, ruta inválida)
    debe surgir como DumpReportError, no como OSError crudo."""

    def test_output_dir_conflicting_with_existing_file_raises_dump_report_error(self, tmp_path: Path):
        # output_dir apunta a un fichero ya existente: mkdir(exist_ok=True) falla con
        # FileExistsError (subclase de OSError) porque no es un directorio.
        blocking_file = tmp_path / "reports"
        blocking_file.write_text("not a directory")
        reporter = DumpReporter(blocking_file)

        with pytest.raises(DumpReportError):
            reporter.generate(table_results=[], missing_tables=[], constraint_results=[], cycles=[])

    def test_wb_save_failure_wrapped_in_dump_report_error(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        import openpyxl as openpyxl_module

        def _boom(self, path, *args, **kwargs):  # noqa: ANN001, ARG001
            raise OSError("disk full")

        monkeypatch.setattr(openpyxl_module.Workbook, "save", _boom)
        reporter = DumpReporter(tmp_path)

        with pytest.raises(DumpReportError, match="disk full"):
            reporter.generate(table_results=[], missing_tables=[], constraint_results=[], cycles=[])


# ---------------------------------------------------------------------------
# Hoja 1: Auditoría Tablas
# ---------------------------------------------------------------------------


class TestAuditSheet:
    def _get_audit_ws(
        self,
        tmp_path: Path,
        results: list[TableDumpResult],
        **kwargs,
    ) -> Worksheet:
        reporter = DumpReporter(tmp_path)
        path = reporter.generate(
            table_results=results,
            missing_tables=kwargs.get("missing_tables", []),
            constraint_results=kwargs.get("constraint_results", []),
            cycles=kwargs.get("cycles", []),
        )
        wb = openpyxl.load_workbook(path)
        return wb["Auditoría Tablas"]

    def test_header_row_has_nine_columns(self, tmp_path: Path):
        ws = self._get_audit_ws(tmp_path, [])
        headers = [ws.cell(1, c).value for c in range(1, 10)]
        assert all(h is not None for h in headers)

    def test_table_result_appears_in_data_row(self, tmp_path: Path):
        ws = self._get_audit_ws(tmp_path, [_table_result("orders", "success", source_rows=42)])
        row2_values = [ws.cell(2, c).value for c in range(1, 8)]
        assert "orders" in row2_values
        assert 42 in row2_values

    def test_missing_target_row_present(self, tmp_path: Path):
        result = _table_result("ghost_table", "missing_target", source_rows=0)
        ws = self._get_audit_ws(tmp_path, [result])
        statuses = [ws.cell(2, 7).value]
        assert "MISSING_TARGET" in statuses

    def test_failed_constraint_appended(self, tmp_path: Path):
        cr = _constraint_result("failed")
        ws = self._get_audit_ws(tmp_path, [_table_result()], constraint_results=[cr])
        # The failed constraint section should be somewhere in the sheet
        all_values = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any(v == "orders" for v in all_values)

    def test_cycle_flag_shown(self, tmp_path: Path):
        result = _table_result("circular_a", in_cycle=True)
        ws = self._get_audit_ws(tmp_path, [result], cycles=[["circular_a"]])
        cycle_col_values = [ws.cell(2, 3).value]
        assert "Sí" in cycle_col_values


# ---------------------------------------------------------------------------
# Hoja 2: Tablas Faltantes
# ---------------------------------------------------------------------------


class TestMissingTablesSheet:
    def _get_ws(self, tmp_path: Path, missing: list[str]) -> Worksheet:
        reporter = DumpReporter(tmp_path)
        path = reporter.generate(
            table_results=[],
            missing_tables=missing,
            constraint_results=[],
            cycles=[],
        )
        wb = openpyxl.load_workbook(path)
        return wb["Tablas Faltantes"]

    def test_empty_missing_shows_placeholder(self, tmp_path: Path):
        ws = self._get_ws(tmp_path, [])
        assert ws.cell(2, 1).value is not None
        assert "sin" in str(ws.cell(2, 1).value).lower()

    def test_missing_table_name_in_row(self, tmp_path: Path):
        ws = self._get_ws(tmp_path, ["orphan_table"])
        row2 = [ws.cell(2, c).value for c in range(1, 4)]
        assert "orphan_table" in row2

    def test_multiple_missing_tables(self, tmp_path: Path):
        ws = self._get_ws(tmp_path, ["t1", "t2", "t3"])
        names = [ws.cell(r, 1).value for r in range(2, 5)]
        assert set(names) == {"t1", "t2", "t3"}


# ---------------------------------------------------------------------------
# Hoja 3: Columnas Inexistentes
# ---------------------------------------------------------------------------


class TestMissingColumnsSheet:
    """La hoja 3 se deriva de ``TableDumpResult.skipped_columns`` (sin comparación de esquemas)."""

    def _get_ws(self, tmp_path: Path, table_results: list[TableDumpResult]) -> Worksheet:
        reporter = DumpReporter(tmp_path)
        path = reporter.generate(
            table_results=table_results,
            missing_tables=[],
            constraint_results=[],
            cycles=[],
        )
        wb = openpyxl.load_workbook(path)
        return wb["Columnas Inexistentes"]

    def test_no_missing_columns_shows_placeholder(self, tmp_path: Path):
        ws = self._get_ws(tmp_path, [_table_result("orders", "success", skipped=[])])
        assert ws.cell(2, 1).value is not None
        assert "sin" in str(ws.cell(2, 1).value).lower()

    def test_missing_column_row_appears(self, tmp_path: Path):
        ws = self._get_ws(tmp_path, [_table_result("orders", "success", skipped=["extra_col"])])
        row2 = [ws.cell(2, c).value for c in range(1, 6)]
        assert "orders" in row2
        assert "extra_col" in row2

    def test_only_skipped_columns_are_reported(self, tmp_path: Path):
        # Una tabla sin skipped_columns no aporta filas; solo las columnas omitidas aparecen.
        ws = self._get_ws(
            tmp_path,
            [
                _table_result("a", "success", skipped=[]),
                _table_result("b", "success", skipped=["only_in_source"]),
            ],
        )
        all_vals = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        assert "only_in_source" in all_vals
        # La tabla 'a' (sin skipped) no genera ninguna fila propia.
        table_col = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "a" not in table_col

    def test_missing_target_table_with_empty_skipped_not_listed(self, tmp_path: Path):
        # Tablas ausentes en destino llegan con skipped vacío → no figuran en la hoja 3.
        ws = self._get_ws(tmp_path, [_table_result("ghost", "missing_target", skipped=[])])
        row2_col1 = ws.cell(2, 1).value
        assert row2_col1 != "ghost"
        assert "sin" in str(row2_col1).lower()


# ---------------------------------------------------------------------------
# Hoja "Auditoría Bulk" (resumen de la auditoría de bulk del volcado)
# ---------------------------------------------------------------------------


def _summary(**overrides: object) -> DumpAuditSummary:
    values: dict[str, object] = {
        "total_operaciones": 4,
        "total_filas": 300,
        "tiempo_total_seg": 1.5,
        "tiempo_promedio_seg": 0.375,
        "tablas_auditadas": 2,
        "tabla_mas_activa": "orders",
        "audit_path": Path("audit.csv"),
        "audited_rows": 300,
    }
    values.update(overrides)
    return DumpAuditSummary(**values)  # type: ignore[arg-type]


def _bulk_audit_sheet(tmp_path: Path, summary: DumpAuditSummary | None) -> openpyxl.Workbook:
    path = DumpReporter(tmp_path).generate(
        table_results=[_table_result()], missing_tables=[], constraint_results=[], cycles=[], audit_summary=summary
    )
    return openpyxl.load_workbook(path)


def _metrics(ws: Worksheet) -> dict[str, object]:
    return {row[0].value: row[1].value for row in ws.iter_rows(min_row=2) if row[0].value and row[1].value is not None}


def _fill(ws: Worksheet, metrica: str) -> str | None:
    for row in ws.iter_rows(min_row=2):
        if row[0].value == metrica:
            color = row[1].fill.fgColor.rgb if row[1].fill.fill_type == "solid" else None
            return str(color) if color else None
    raise AssertionError(metrica)


class TestBulkAuditSheet:
    def test_sheet_is_added_last_with_the_summary_values(self, tmp_path: Path) -> None:
        wb = _bulk_audit_sheet(tmp_path, _summary(audit_failures=1, unaudited_committed_rows=2))

        assert wb.sheetnames[-1] == "Auditoría Bulk"
        assert wb.sheetnames[:3] == ["Auditoría Tablas", "Tablas Faltantes", "Columnas Inexistentes"]
        metrics = _metrics(wb["Auditoría Bulk"])
        assert metrics["CSV de auditoría"] == "audit.csv"
        assert metrics["Operaciones en el CSV"] == 4
        assert metrics["Filas en el CSV"] == 300
        assert metrics["Filas escritas por el volcado"] == 300
        assert metrics["Tabla más activa"] == "orders"
        assert metrics["Registros no escritos"] == 1
        assert metrics["Filas confirmadas sin auditar"] == 2
        assert metrics["Registros sin datos en destino"] == 0

    def test_sheet_goes_after_the_indexes_sheet(self, tmp_path: Path) -> None:
        # Añadir la hoja al final no cambia la posición de las existentes (quien lea wb.worksheets[i]).
        path = DumpReporter(tmp_path).generate(
            table_results=[],
            missing_tables=[],
            constraint_results=[],
            cycles=[],
            index_results=[],
            audit_summary=_summary(),
        )

        assert openpyxl.load_workbook(path).sheetnames[-2:] == ["Índices", "Auditoría Bulk"]

    def test_without_summary_there_is_no_sheet(self, tmp_path: Path) -> None:
        assert "Auditoría Bulk" not in _bulk_audit_sheet(tmp_path, None).sheetnames

    def test_incidents_are_highlighted_and_zero_counters_are_not(self, tmp_path: Path) -> None:
        ws = _bulk_audit_sheet(tmp_path, _summary(audit_failures=3, audit_uncommitted_records=1))["Auditoría Bulk"]

        assert _fill(ws, "Registros no escritos") == DumpReporter._COL_RED
        assert _fill(ws, "Registros sin datos en destino") == DumpReporter._COL_RED
        assert _fill(ws, "Filas confirmadas sin auditar") is None
        assert _fill(ws, "Filas en el CSV") is None  # CSV y contadores coinciden

    def test_csv_and_counters_mismatch_is_highlighted(self, tmp_path: Path) -> None:
        ws = _bulk_audit_sheet(tmp_path, _summary(total_filas=500, audited_rows=300))["Auditoría Bulk"]

        assert _fill(ws, "Filas en el CSV") == DumpReporter._COL_YELLOW
        assert _fill(ws, "Filas escritas por el volcado") == DumpReporter._COL_YELLOW

    def test_table_name_from_the_csv_is_stored_as_text_not_as_a_formula(self, tmp_path: Path) -> None:
        wb = _bulk_audit_sheet(tmp_path, _summary(tabla_mas_activa='=HYPERLINK("http://x","clic")'))

        cell = next(r[1] for r in wb["Auditoría Bulk"].iter_rows(min_row=2) if r[0].value == "Tabla más activa")
        assert cell.data_type == "s"
        assert cell.value == '=HYPERLINK("http://x","clic")'

    def test_formulas_are_neutralized_in_every_sheet(self, tmp_path: Path) -> None:
        # Nombres del catálogo de la BD y mensajes de error también llegan al informe sin control.
        path = DumpReporter(tmp_path).generate(
            table_results=[_table_result(name="=1+1", status="error", error_msg="=cmd|' /C calc'!A0")],
            missing_tables=["=SUM(A1)"],
            constraint_results=[],
            cycles=[],
        )

        wb = openpyxl.load_workbook(path)
        formulas = [c.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]
        assert formulas == []


# ---------------------------------------------------------------------------
# Hoja "Secuencias PK" (resincronización de contadores de PK en destino)
# ---------------------------------------------------------------------------


def _pk(status: str = "synced", max_value: int | None = 100, error_msg: str = "") -> PKSyncResult:
    return PKSyncResult(
        table_name="orders",
        column_name="id",
        engine_dialect="postgresql",
        max_value=max_value,
        status=status,  # type: ignore[arg-type]
        error_msg=error_msg,
    )


def _pk_workbook(tmp_path: Path, pk_sync_results: list[PKSyncResult] | None, **kwargs: object) -> openpyxl.Workbook:
    path = DumpReporter(tmp_path).generate(
        table_results=[],
        missing_tables=[],
        constraint_results=[],
        cycles=[],
        pk_sync_results=pk_sync_results,
        **kwargs,  # type: ignore[arg-type]
    )
    return openpyxl.load_workbook(path)


class TestPKSyncSheet:
    def test_one_row_per_column_with_its_values(self, tmp_path: Path) -> None:
        ws = _pk_workbook(tmp_path, [_pk(), _pk(status="no_rows", max_value=None)])["Secuencias PK"]

        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        assert rows == [
            ["orders", "id", "postgresql", 100, "SYNCED", "—"],
            ["orders", "id", "postgresql", "—", "NO_ROWS", "—"],
        ]

    def test_status_colors(self, tmp_path: Path) -> None:
        results = [_pk(), _pk(status="error", error_msg="permiso denegado"), _pk(status="unsupported")]
        ws = _pk_workbook(tmp_path, results)["Secuencias PK"]

        colors = [ws.cell(row=r, column=5).fill.fgColor.rgb for r in (2, 3, 4)]
        assert colors == [DumpReporter._COL_GREEN, DumpReporter._COL_RED, DumpReporter._COL_ORANGE]
        assert ws.cell(row=3, column=6).value == "permiso denegado"

    def test_without_results_there_is_no_sheet(self, tmp_path: Path) -> None:
        assert "Secuencias PK" not in _pk_workbook(tmp_path, None).sheetnames

    def test_empty_list_shows_a_message(self, tmp_path: Path) -> None:
        ws = _pk_workbook(tmp_path, [])["Secuencias PK"]

        assert ws.cell(row=2, column=1).value == "Sin contadores de PK que sincronizar"

    def test_sheet_goes_last_after_the_other_optional_sheets(self, tmp_path: Path) -> None:
        wb = _pk_workbook(tmp_path, [_pk()], index_results=[], audit_summary=_summary())

        assert wb.sheetnames[-3:] == ["Índices", "Auditoría Bulk", "Secuencias PK"]
