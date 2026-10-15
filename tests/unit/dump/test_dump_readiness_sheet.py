"""Hoja «Condiciones del motor» del informe de dump: origen y destino (EngineReadinessAdvisor)."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from sqlalchemy import create_engine, text
from sqlalchemy.engine import Engine

from util_bbdd import DumpOrchestrator
from util_bbdd.connection.performance_advisor import EngineReadinessAdvisor
from util_bbdd.dump import DumpConfig
from util_bbdd.dump.reporters import DumpReporter

SHEET = "Condiciones del motor"


def _sqlite(path: Path) -> Engine:
    engine = create_engine(f"sqlite:///{path}")
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE IF NOT EXISTS t (id INTEGER PRIMARY KEY, n TEXT)"))
    return engine


@pytest.fixture
def engines(tmp_path: Path) -> Iterator[tuple[Engine, Engine]]:
    src, dst = _sqlite(tmp_path / "src.db"), _sqlite(tmp_path / "dst.db")
    with src.begin() as conn:
        conn.execute(text("INSERT INTO t VALUES (1, 'a'), (2, 'b')"))
    yield src, dst
    src.dispose()
    dst.dispose()


def _first_column(report: Path) -> list[object]:
    return [row[0].value for row in openpyxl.load_workbook(report)[SHEET].iter_rows()]


def test_dump_report_ends_with_one_block_per_engine(engines: tuple[Engine, Engine], tmp_path: Path) -> None:
    src, dst = engines
    result = DumpOrchestrator(DumpConfig(source_engine=src, target_engine=dst, output_dir=tmp_path / "out")).run()

    assert openpyxl.load_workbook(result.report).sheetnames[-1] == SHEET
    column = _first_column(result.report)
    assert f"Origen · {EngineReadinessAdvisor.describe_engine(src)}" in column
    assert f"Destino · {EngineReadinessAdvisor.describe_engine(dst)}" in column


def test_same_engine_for_source_and_target_gives_a_single_block(tmp_path: Path) -> None:
    engine = _sqlite(tmp_path / "same.db")
    orchestrator = DumpOrchestrator(DumpConfig(source_engine=engine, target_engine=engine, output_dir=tmp_path))

    blocks = orchestrator._engine_readiness()

    assert [title for title, _ in blocks] == [f"Origen y destino · {EngineReadinessAdvisor.describe_engine(engine)}"]
    engine.dispose()


def test_a_failed_verification_does_not_stop_the_dump(engines: tuple[Engine, Engine], tmp_path: Path) -> None:
    src, dst = engines
    with patch.object(EngineReadinessAdvisor, "engine_report", side_effect=RuntimeError("boom")):
        result = DumpOrchestrator(DumpConfig(source_engine=src, target_engine=dst, output_dir=tmp_path / "out")).run()

    column = _first_column(result.report)
    assert column.count("⚠️ No se pudieron verificar las condiciones de este motor.") == 2
    assert {"Origen", "Destino"} <= set(column)


def test_reporter_without_engine_readiness_keeps_the_previous_sheets(tmp_path: Path) -> None:
    path = DumpReporter(tmp_path).generate(table_results=[], missing_tables=[], constraint_results=[], cycles=[])

    assert SHEET not in openpyxl.load_workbook(path).sheetnames
