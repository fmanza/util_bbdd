"""Plantilla Excel de la documentación y estilo de tabla propio del Excel (util_bbdd_web G40, 4.1.0)."""

from __future__ import annotations

import re
import sqlite3

from pathlib import Path
from typing import Any

import openpyxl
import pytest

from openpyxl.writer.theme import theme_xml

from util_bbdd import SyncDatabaseManager
from util_bbdd.doc import DatabaseDocumentation
from util_bbdd.doc.infrastructure.renderers.excel_template import ExcelTemplate
from util_bbdd.exceptions import DocumentationConfigError

pytestmark = pytest.mark.unit

IMPRIMIR: dict[str, Any] = {
    "tablas": dict.fromkeys(("imprimir_tablas", "columnas", "pk", "fk", "indices"), True),
    "vistas": {"imprimir_vistas": True, "columnas": True},
}
FIXED = set(ExcelTemplate.FIXED_COLORS)


@pytest.fixture
def sqlite_db(tmp_path: Path) -> Path:
    path = tmp_path / "demo.db"
    con = sqlite3.connect(path)
    con.executescript(
        """
        CREATE TABLE customers (id INTEGER PRIMARY KEY, name TEXT);
        CREATE TABLE orders (id INTEGER PRIMARY KEY, customer_id INTEGER REFERENCES customers(id));
        CREATE VIEW v_orders AS SELECT id FROM orders;
        """
    )
    con.close()
    return path


def _template(tmp_path: Path, *, theme: str | None = None) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.oddHeader.center.text = "ACME &P"
    sheet.oddFooter.right.text = "Confidencial&G"
    sheet.page_setup.orientation = "landscape"
    workbook.loaded_theme = (
        theme or re.sub(r'(<a:minorFont>\s*<a:latin typeface=")Calibri"', r'\1Arial"', theme_xml)
    ).encode("utf-8")
    path = tmp_path / "plantilla.xlsx"
    workbook.save(path)
    return path


def _generate(sqlite_db: Path, tmp_path: Path, **imprimir: Any) -> openpyxl.Workbook:
    with SyncDatabaseManager("sqlite", {"database": str(sqlite_db)}, use_pool=False) as manager:
        result = DatabaseDocumentation(
            manager,
            directorio_salida=tmp_path / "out",
            fichero_salida=Path("demo"),
            imprimir={**IMPRIMIR, **imprimir},
            formatos=["excel"],
        ).generate()
    return openpyxl.load_workbook(result.outputs["excel"])


def _fills(workbook: openpyxl.Workbook) -> list[Any]:
    return [
        cell.fill.fgColor
        for ws in workbook.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.fill is not None and cell.fill.fill_type == "solid"
    ]


def test_without_template_the_workbook_keeps_its_fixed_colors(sqlite_db: Path, tmp_path: Path) -> None:
    fills = _fills(_generate(sqlite_db, tmp_path))

    assert fills
    assert {color.rgb[-6:] for color in fills if color.type == "rgb"} & FIXED


def test_template_brings_theme_fonts_header_and_page(sqlite_db: Path, tmp_path: Path) -> None:
    workbook = _generate(sqlite_db, tmp_path, plantilla_excel=_template(tmp_path))

    fills = _fills(workbook)
    assert not {color.rgb[-6:] for color in fills if color.type == "rgb"} & FIXED  # ningún color fijo
    assert any(color.type == "theme" for color in fills)
    assert b'typeface="Arial"' in workbook.loaded_theme
    named_fonts = {
        cell.font.name for ws in workbook.worksheets for row in ws.iter_rows() for cell in row if cell.font.bold
    }
    assert "Arial" in named_fonts
    for ws in workbook.worksheets:
        assert ws.oddHeader.center.text == "ACME &P"
        assert ws.oddFooter.right.text == "Confidencial"  # sin la imagen (&G)
        assert ws.page_setup.orientation == "landscape"


def test_table_style_excel_applies_to_every_table(sqlite_db: Path, tmp_path: Path) -> None:
    workbook = _generate(sqlite_db, tmp_path, table_style="Light List", table_style_excel="TableStyleLight9")

    styles = {table.tableStyleInfo.name for ws in workbook.worksheets for table in ws.tables.values()}
    assert styles == {"TableStyleLight9"}


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"no soy un libro", "no es un libro"),
        (None, "referencia imágenes"),
    ],
)
def test_invalid_template_is_rejected_when_the_documentation_is_built(
    sqlite_db: Path, tmp_path: Path, content: bytes | None, message: str
) -> None:
    if content is None:
        path = _template(tmp_path, theme=theme_xml.replace("<a:theme ", '<a:theme r:embed="rId1" ', 1))
    else:
        path = tmp_path / "plantilla.xlsx"
        path.write_bytes(content)

    with SyncDatabaseManager("sqlite", {"database": str(sqlite_db)}, use_pool=False) as manager:
        with pytest.raises(DocumentationConfigError, match=message):
            DatabaseDocumentation(manager, imprimir={"plantilla_excel": path}, formatos=["excel"])


def test_missing_template_is_rejected_only_when_excel_is_requested(sqlite_db: Path, tmp_path: Path) -> None:
    missing = tmp_path / "no_existe.xlsx"
    with SyncDatabaseManager("sqlite", {"database": str(sqlite_db)}, use_pool=False) as manager:
        with pytest.raises(DocumentationConfigError, match="no encontrada"):
            DatabaseDocumentation(manager, imprimir={"plantilla_excel": missing}, formatos=["excel"])
        DatabaseDocumentation(manager, imprimir={"plantilla_excel": missing}, formatos=["word"])
