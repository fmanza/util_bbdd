"""Estilos de tabla de Word (resueltos una vez por documento, P-02) y de Excel (sin aviso por defecto, A-11)."""

from __future__ import annotations

import logging

from typing import Any

import docx
import pytest

from util_bbdd.doc.infrastructure.renderers.excel_generator import ExcelGenerator
from util_bbdd.doc.infrastructure.renderers.word_generator import WordGenerator


def _count_style_scans(monkeypatch: pytest.MonkeyPatch, generator: WordGenerator) -> list[int]:
    scans: list[int] = []
    original = generator._resolve_table_style

    def _counting(*args: Any) -> Any:
        scans.append(1)
        return original(*args)

    monkeypatch.setattr(generator, "_resolve_table_style", _counting)
    return scans


def test_word_table_style_is_resolved_once_per_document(monkeypatch: pytest.MonkeyPatch) -> None:
    generator = WordGenerator(database="demo", headers={}, table_style="Table Grid")
    scans = _count_style_scans(monkeypatch, generator)

    first = docx.Document()
    for _ in range(5):
        generator._add_table(first, "T", ["a", "b"], [["1", "2"]])
    second = docx.Document()
    generator._add_table(second, "T", ["a"], [["1"]])

    assert len(scans) == 2  # una por documento
    assert all(table.style.name == "Table Grid" for table in first.tables)
    assert second.tables[0].style.name == "Table Grid"


def test_word_unknown_table_style_falls_back_once(monkeypatch: pytest.MonkeyPatch, caplog: Any) -> None:
    generator = WordGenerator(database="demo", headers={}, table_style="No existe")
    document = docx.Document()

    with caplog.at_level(logging.WARNING):
        for _ in range(3):
            generator._add_table(document, "T", ["a"], [["1"]])

    assert [table.style.name for table in document.tables] == ["Table Grid"] * 3
    assert sum("no disponible" in record.getMessage() for record in caplog.records) == 1


def test_excel_accepts_the_default_word_style_without_warning(caplog: Any) -> None:
    with caplog.at_level(logging.WARNING):
        generator = ExcelGenerator(database="demo", headers={}, table_style="Table Grid")

    assert generator.table_style == "TableStyleMedium2"
    assert not caplog.records
