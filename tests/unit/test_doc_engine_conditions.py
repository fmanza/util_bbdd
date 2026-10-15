"""Sección «Condiciones del motor» de DatabaseDocumentation (EngineReadinessAdvisor) en Word, Excel y HTML."""

from __future__ import annotations

import sqlite3

from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pytest

from util_bbdd import SyncDatabaseManager
from util_bbdd.connection import AdvisoryFinding, EngineReadinessAdvisor, EngineReadinessReport
from util_bbdd.doc import DatabaseDocumentation
from util_bbdd.doc.application.use_cases.engine_conditions import EngineConditionsSectionBuilder
from util_bbdd.doc.infrastructure.adapters.sync_database_manager_adapter import SyncDatabaseManagerAdapter
from util_bbdd.doc.infrastructure.renderers.shared.print_config_manager import PrintConfigManager

IMPRIMIR = {"tablas": {"imprimir_tablas": True}, "condiciones_motor": {"imprimir_condiciones_motor": True}}
SECRET = "s3creto-centinela"


def _report(*findings: AdvisoryFinding, dynamic: bool = True, supported: bool = True, error: str | None = None):
    return EngineReadinessReport(
        db_type="mysql",
        findings=findings,
        checked_at=datetime(2026, 10, 15, 14, 0, 0, tzinfo=UTC),
        dynamic=dynamic,
        supported=supported,
        error=error,
    )


class _FailingEngine:
    """Engine cuya conexión falla con un mensaje que lleva credenciales (como hacen algunos drivers)."""

    def connect(self) -> Any:
        raise OSError(f"mysql://root:{SECRET}@db:3306 inalcanzable")


# --- EngineReadinessAdvisor.assess ------------------------------------------------------------------------------


def test_assess_without_engine_runs_static_checks_only() -> None:
    report = EngineReadinessAdvisor().assess("mysql", {"pool_size": 2})

    assert report.supported is True
    assert report.dynamic is False
    assert report.error == "Sin engine: solo checks estáticos"
    assert [(f.level, f.param) for f in report.findings] == [("WARNING", "pool_size")]
    assert report.checked_at.tzinfo is not None


def test_assess_marks_engines_without_checks_as_unsupported() -> None:
    report = EngineReadinessAdvisor().assess("hsqldb", {"pool_size": 2})

    assert (report.supported, report.dynamic, report.findings, report.error) == (False, False, (), None)


def test_assess_falls_back_to_static_checks_and_hides_the_connection_error_message() -> None:
    report = EngineReadinessAdvisor().assess("mysql", {"pool_size": 2}, engine=_FailingEngine())  # type: ignore[arg-type]

    assert report.dynamic is False
    assert report.error == "Sin conexión para los checks dinámicos (OSError)"
    assert SECRET not in repr(report)
    assert [f.param for f in report.findings] == ["pool_size"]


def test_manager_get_engine_readiness_runs_dynamic_checks_on_sqlite() -> None:
    with SyncDatabaseManager("sqlite", {"database": ":memory:"}) as manager:
        report = manager.get_engine_readiness()

    assert (report.db_type, report.supported, report.dynamic, report.error) == ("sqlite", True, True, None)
    assert report.findings


def test_manager_get_engine_readiness_without_engine_keeps_static_checks(monkeypatch: pytest.MonkeyPatch) -> None:
    def _no_engine(self: SyncDatabaseManager) -> Any:
        raise RuntimeError(SECRET)

    monkeypatch.setattr(SyncDatabaseManager, "get_sqlalchemy_engine", _no_engine)
    with SyncDatabaseManager("sqlite", {"database": ":memory:"}) as manager:
        report = manager.get_engine_readiness()

    assert report.supported is True
    assert report.dynamic is False
    assert report.error == "Sin engine para los checks dinámicos (RuntimeError)"


# --- EngineConditionsSectionBuilder -----------------------------------------------------------------------------


def test_builder_orders_rows_by_severity_then_parameter_and_counts_every_level() -> None:
    section = EngineConditionsSectionBuilder().build(
        _report(
            AdvisoryFinding("OK", "b", "1", ""),
            AdvisoryFinding("WARNING", "z", "1", "subir"),
            AdvisoryFinding("CRITICAL", "y", "0", "corregir"),
            AdvisoryFinding("WARNING", "a", "1", "subir"),
            AdvisoryFinding("UNKNOWN", "q", "?", ""),  # type: ignore[arg-type]
        )
    )

    assert [row[:2] for row in section.rows] == [
        ("CRITICAL", "y"),
        ("WARNING", "a"),
        ("WARNING", "z"),
        ("OK", "b"),
        ("UNKNOWN", "q"),
    ]
    assert section.summary == (("CRITICAL", 1), ("WARNING", 2), ("INFO", 0), ("OK", 1), ("UNKNOWN", 1))
    assert section.metadata == (
        ("Motor", "mysql"),
        ("Verificado (UTC)", "2026-10-15 14:00:00"),
        ("Checks dinámicos", "Sí"),
    )
    assert section.notices == ()
    assert section.headers == EngineConditionsSectionBuilder.DEFAULT_HEADERS


def test_builder_converts_the_verification_time_to_utc() -> None:
    report = EngineReadinessReport(
        db_type="mysql",
        findings=(),
        checked_at=datetime(2026, 10, 15, 16, 0, 0, tzinfo=timezone(timedelta(hours=2))),
        dynamic=True,
        supported=True,
    )

    assert dict(EngineConditionsSectionBuilder().build(report).metadata)["Verificado (UTC)"] == "2026-10-15 14:00:00"


@pytest.mark.parametrize(
    ("report", "expected"),
    [
        (None, [EngineConditionsSectionBuilder.NOTICE_UNAVAILABLE]),
        (
            _report(supported=False, dynamic=False),
            ["El motor 'mysql' no tiene verificaciones de condiciones de ejecución."],
        ),
        (_report(), [EngineConditionsSectionBuilder.NOTICE_NO_FINDINGS]),
        (
            _report(AdvisoryFinding("OK", "a", "1", ""), dynamic=False, error="Sin engine: solo checks estáticos"),
            [
                "Solo checks estáticos (configuración de conexión), sin variables del servidor: Sin engine: solo checks estáticos."
            ],
        ),
    ],
    ids=["origen-sin-verificacion", "motor-sin-checks", "sin-hallazgos", "solo-estaticos"],
)
def test_builder_explains_empty_or_partial_sections(report: EngineReadinessReport | None, expected: list[str]) -> None:
    assert list(EngineConditionsSectionBuilder().build(report).notices) == expected


def test_builder_uses_custom_headers_and_providers_without_the_port_get_the_notice() -> None:
    section = EngineConditionsSectionBuilder(["L", "P", "V", "R"]).build(_report())

    assert section.headers == ("L", "P", "V", "R")
    assert EngineConditionsSectionBuilder.of(object()).notices == (EngineConditionsSectionBuilder.NOTICE_UNAVAILABLE,)


@pytest.mark.parametrize("headers", [["Nivel", "Parámetro"], ["a", "b", "c", "d", "e"], [], None])
def test_builder_falls_back_to_default_headers_when_they_do_not_match_the_rows(headers: list[str] | None) -> None:
    section = EngineConditionsSectionBuilder(headers).build(_report(AdvisoryFinding("OK", "a", "1", "")))

    assert section.headers == EngineConditionsSectionBuilder.DEFAULT_HEADERS


def test_adapter_returns_none_when_the_source_cannot_verify_the_engine() -> None:
    class _Source:
        db_type = "sqlite"
        database = "x"
        conn_params: dict[str, Any] = {}

    assert SyncDatabaseManagerAdapter(_Source(), {}).get_engine_readiness() is None  # type: ignore[arg-type]


# --- PrintConfigManager ------------------------------------------------------------------------------------------


def test_engine_conditions_section_is_disabled_by_default() -> None:
    for config in (PrintConfigManager.normalize(None), PrintConfigManager.normalize({})):
        assert config["condiciones_motor"] == {"imprimir_condiciones_motor": False}
        assert PrintConfigManager.resolve_for_word(config)["include_engine_conditions"] is False
        assert PrintConfigManager.resolve_for_excel(config)["include_engine_conditions"] is False


def test_engine_conditions_flag_reaches_every_format() -> None:
    config = PrintConfigManager.normalize({"imprimir": {"condiciones_motor": {"imprimir_condiciones_motor": 1}}})

    assert config["condiciones_motor"] == {"imprimir_condiciones_motor": True}
    assert PrintConfigManager.normalize(config)["condiciones_motor"] == {"imprimir_condiciones_motor": True}
    assert PrintConfigManager.resolve_for_word(config)["include_engine_conditions"] is True
    assert PrintConfigManager.resolve_for_excel(config)["include_engine_conditions"] is True
    html = PrintConfigManager.resolve_for_html(config, output_dir=Path("out"), default_template_file=Path("d.j2"))
    assert html["include_engine_conditions"] is True


# --- DatabaseDocumentation de extremo a extremo -----------------------------------------------------------------


@pytest.fixture
def sqlite_db(tmp_path: Path) -> Path:
    path = tmp_path / "demo.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE customers (id INTEGER PRIMARY KEY, email TEXT)")
    con.close()
    return path


def _generate(sqlite_db: Path, out: Path, imprimir: dict[str, Any]) -> Any:
    with SyncDatabaseManager("sqlite", {"database": str(sqlite_db)}, use_pool=False) as manager:
        return DatabaseDocumentation(
            manager,
            directorio_salida=out,
            fichero_salida=Path("demo"),
            imprimir=imprimir,
            formatos=["word", "excel", "html"],
        ).generate()


def _word_text(path: Path) -> str:
    document = docx.Document(str(path))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def _all_html(html_dir: Path) -> str:
    return "\n".join(page.read_text(encoding="utf-8") for page in html_dir.rglob("*.html"))


def test_section_appears_in_word_excel_and_html(sqlite_db: Path, tmp_path: Path) -> None:
    result = _generate(sqlite_db, tmp_path / "out", IMPRIMIR)

    word = _word_text(result.outputs["word"])
    assert "Condiciones del motor" in word
    assert "Resumen por severidad" in word
    assert "journal_mode" in word or "timeout" in word

    sheet = openpyxl.load_workbook(result.outputs["excel"])["Condiciones del motor"]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert ["Nivel", "Parámetro", "Valor actual", "Recomendación"] in values
    assert ["Checks dinámicos", "Sí", None, None] in values

    html_dir = result.outputs["html"]
    page = (html_dir / "condiciones" / "condiciones-motor.html").read_text(encoding="utf-8")
    assert "Verificación de condiciones de ejecución" in page
    assert "<h1>Condiciones del motor: sqlite</h1>" in page
    assert "Total de filas recopiladas" not in page
    assert 'href="condiciones/condiciones-motor.html"' in (html_dir / "index.html").read_text(encoding="utf-8")


def test_section_is_absent_unless_requested(sqlite_db: Path, tmp_path: Path) -> None:
    result = _generate(sqlite_db, tmp_path / "out", {"tablas": {"imprimir_tablas": True}})

    assert "Condiciones del motor" not in _word_text(result.outputs["word"])
    assert "Condiciones del motor" not in openpyxl.load_workbook(result.outputs["excel"]).sheetnames
    assert not (result.outputs["html"] / "condiciones").exists()
    index = (result.outputs["html"] / "index.html").read_text(encoding="utf-8")
    assert 'value="engine-conditions"' not in index
    assert "scope-engine-conditions-block" not in index


def test_engine_is_verified_once_for_all_formats(
    sqlite_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls: list[str] = []
    original = SyncDatabaseManager.get_engine_readiness

    def _counting(self: SyncDatabaseManager) -> EngineReadinessReport:
        calls.append(self.db_type)
        return original(self)

    monkeypatch.setattr(SyncDatabaseManager, "get_engine_readiness", _counting)
    _generate(sqlite_db, tmp_path / "out", IMPRIMIR)

    assert calls == ["sqlite"]


def test_failed_verification_does_not_stop_generation_nor_leak_its_message(
    sqlite_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _failing(self: SyncDatabaseManager) -> EngineReadinessReport:
        raise RuntimeError(f"password={SECRET}")

    monkeypatch.setattr(SyncDatabaseManager, "get_engine_readiness", _failing)
    result = _generate(sqlite_db, tmp_path / "out", IMPRIMIR)

    notice = "No se pudieron verificar las condiciones del motor (RuntimeError)."
    word = _word_text(result.outputs["word"])
    assert notice in word
    excel = openpyxl.load_workbook(result.outputs["excel"])
    excel_text = "\n".join(str(c.value) for ws in excel.worksheets for row in ws.iter_rows() for c in row)
    html_text = _all_html(result.outputs["html"])
    assert notice in html_text
    for text in (word, excel_text, html_text):
        assert SECRET not in text
