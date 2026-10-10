"""DatabaseDocumentation de extremo a extremo sobre SQLite: Word, Excel y HTML con todas las secciones."""

from __future__ import annotations

import html
import re
import sqlite3

from pathlib import Path

import docx
import openpyxl
import pytest

from util_bbdd import SyncDatabaseManager
from util_bbdd.doc import DatabaseDocumentation

IMPRIMIR = {
    "tablas": dict.fromkeys(("imprimir_tablas", "ddl", "columnas", "pk", "fk", "unicas", "indices"), True),
    "vistas": {"imprimir_vistas": True, "ddl": True, "columnas": True},
    "rutinas": {"imprimir_procedures": True, "imprimir_function": True},
    "almacenamiento": {"imprimir_almacenamiento": True},
    "info_database": {"imprimir_info_database": True, "version": True, "charset_collation": True},
}


@pytest.fixture
def sqlite_db(tmp_path: Path) -> Path:
    path = tmp_path / "demo.db"
    con = sqlite3.connect(path)
    con.executescript(
        """
        CREATE TABLE customers (id INTEGER PRIMARY KEY, email TEXT NOT NULL UNIQUE, name TEXT DEFAULT 'x');
        CREATE TABLE orders (
            id INTEGER PRIMARY KEY,
            customer_id INTEGER NOT NULL REFERENCES customers(id) ON DELETE CASCADE,
            total NUMERIC(10, 2)
        );
        CREATE INDEX ix_orders_customer ON orders (customer_id);
        CREATE VIEW v_orders AS SELECT o.id, c.email FROM orders o JOIN customers c ON c.id = o.customer_id;
        CREATE VIEW v_literals AS SELECT id, CASE WHEN id < 2 THEN '&amp;<b>' ELSE '>' END AS c FROM orders;
        """
    )
    con.close()
    return path


def test_generates_word_excel_and_html_from_sqlite(sqlite_db: Path, tmp_path: Path) -> None:
    with SyncDatabaseManager("sqlite", {"database": str(sqlite_db)}, use_pool=False) as manager:
        result = DatabaseDocumentation(
            manager,
            directorio_salida=tmp_path / "out",
            fichero_salida=Path("demo"),
            imprimir=IMPRIMIR,
            formatos=["word", "excel", "html"],
        ).generate()

    assert (result.tables, result.views, result.routines) == (2, 2, 0)

    workbook = openpyxl.load_workbook(result.outputs["excel"])
    cells = {str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row}
    assert {"customers", "orders", "Vista: v_orders", "customer_id"} <= cells
    assert any("ix_orders_customer" in value for value in cells)

    word_text = "\n".join(p.text for p in docx.Document(str(result.outputs["word"])).paragraphs)
    assert "Tabla: orders" in word_text
    word_cells = {c.text for t in docx.Document(str(result.outputs["word"])).tables for r in t.rows for c in r.cells}
    assert {"customer_id", "customers", "CASCADE"} <= word_cells

    html_dir = result.outputs["html"]
    assert (html_dir / "index.html").is_file()
    assert any("customer_id" in page.read_text(encoding="utf-8") for page in (html_dir / "tables").glob("*.html"))


def test_html_shows_the_view_definition_as_the_engine_stores_it(sqlite_db: Path, tmp_path: Path) -> None:
    """S-02 (v0.41.1): el HTML decodificaba entidades del DDL (hasta 3 veces) y mostraba ``'&<b>'`` donde el motor guarda
    ``'&amp;<b>'``; ningún motor devuelve DDL con entidades. Word ya lo mostraba bien."""
    imprimir = {"vistas": {"imprimir_vistas": True, "ddl": True}}
    with SyncDatabaseManager("sqlite", {"database": str(sqlite_db)}, use_pool=False) as manager:
        result = DatabaseDocumentation(
            manager, directorio_salida=tmp_path / "out", imprimir=imprimir, formatos=["html", "word"]
        ).generate()

    page = next((result.outputs["html"] / "views").glob("v-literals*.html")).read_text(encoding="utf-8")
    page = re.sub(r"<(script|style)\b.*?</\1>", " ", page, flags=re.IGNORECASE | re.DOTALL)
    shown = html.unescape(re.sub(r"<[^>]+>", " ", page))  # lo que muestra el navegador (sin el JS de resaltado)
    assert "THEN '&amp;<b>' ELSE '>'" in shown
    word = "\n".join(p.text for p in docx.Document(str(result.outputs["word"])).paragraphs)
    assert "THEN '&amp;<b>' ELSE '>'" in word
