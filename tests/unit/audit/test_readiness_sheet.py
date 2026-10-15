"""Hoja «Condiciones del motor» de los Excel de auditoría (audit/readiness.py), con informes falsos."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime, timedelta, timezone

import openpyxl
import pytest

from openpyxl.worksheet.worksheet import Worksheet

from util_bbdd._core.export_safety import SafeWorkbook
from util_bbdd.audit.readiness import ReadinessReportView, ReadinessRows, ReadinessSheetWriter


@dataclass(frozen=True)
class _Finding:
    level: str
    param: str
    current: str
    recommendation: str


@dataclass(frozen=True)
class _Report:
    db_type: str = "mysql"
    findings: tuple[_Finding, ...] = ()
    checked_at: datetime = datetime(2026, 10, 15, 14, 0, tzinfo=UTC)
    dynamic: bool = True
    supported: bool = True
    error: str | None = None


def _cells(ws: Worksheet) -> list[list[object]]:
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def test_rows_are_ordered_by_severity_then_parameter_and_every_level_is_counted() -> None:
    content = ReadinessRows.from_report(
        _Report(
            findings=(
                _Finding("OK", "b", "1", ""),
                _Finding("WARNING", "z", "1", "subir"),
                _Finding("CRITICAL", "y", "0", "corregir"),
                _Finding("WARNING", "a", "1", "subir"),
                _Finding("RARO", "q", "?", ""),
            )
        )
    )

    assert [row[:2] for row in content.rows] == [
        ("CRITICAL", "y"),
        ("WARNING", "a"),
        ("WARNING", "z"),
        ("OK", "b"),
        ("RARO", "q"),
    ]
    assert content.summary == (("CRITICAL", 1), ("WARNING", 2), ("INFO", 0), ("OK", 1), ("RARO", 1))
    assert content.metadata == (
        ("Motor", "mysql"),
        ("Verificado (UTC)", "2026-10-15 14:00:00"),
        ("Checks dinámicos", "Sí"),
    )
    assert content.notices == ()


def test_verification_time_is_shown_in_utc() -> None:
    report = _Report(checked_at=datetime(2026, 10, 15, 16, 0, tzinfo=timezone(timedelta(hours=2))))

    assert dict(ReadinessRows.from_report(report).metadata)["Verificado (UTC)"] == "2026-10-15 14:00:00"


@pytest.mark.parametrize(
    ("report", "notices"),
    [
        (None, [ReadinessRows.NOTICE_UNAVAILABLE]),
        (
            _Report(supported=False, dynamic=False),
            ["El motor 'mysql' no tiene verificaciones de condiciones de ejecución."],
        ),
        (_Report(), [ReadinessRows.NOTICE_NO_FINDINGS]),
        (
            _Report(findings=(_Finding("OK", "a", "1", ""),), dynamic=False, error="Sin conexión (OSError)"),
            ["Solo checks estáticos (configuración de conexión), sin variables del servidor: Sin conexión (OSError)."],
        ),
    ],
    ids=["no-verificado", "motor-sin-checks", "sin-hallazgos", "solo-estaticos"],
)
def test_notices_explain_empty_or_partial_blocks(report: _Report | None, notices: list[str]) -> None:
    assert list(ReadinessRows.from_report(report).notices) == notices


def test_writer_puts_one_block_per_engine_and_neutralizes_formulas() -> None:
    wb = SafeWorkbook()
    ws = wb.create_sheet("Condiciones del motor")
    origin = _Report(findings=(_Finding("WARNING", "pool_size", '=HYPERLINK("x")', "subir"),))
    ReadinessSheetWriter().write(ws, [("Origen · mysql · db:3306/app", origin), ("Destino", None)])

    cells = _cells(ws)
    first_column = [row[0] for row in cells]
    assert first_column[0] == "Condiciones del motor"
    assert "Origen · mysql · db:3306/app" in first_column
    assert "Destino" in first_column
    assert f"⚠️ {ReadinessRows.NOTICE_UNAVAILABLE}" in first_column
    assert ["Nivel", "Parámetro", "Valor actual", "Recomendación"] in cells
    formula_cell = next(cell for row in ws.iter_rows() for cell in row if cell.value == '=HYPERLINK("x")')
    assert formula_cell.data_type == "s"


def test_writer_works_on_a_plain_openpyxl_workbook() -> None:
    """El resumen de bulk escribe con pd.ExcelWriter (libro openpyxl normal, sin SafeWorkbook)."""
    ws = openpyxl.Workbook().create_sheet("Condiciones_motor")
    ReadinessSheetWriter().write(ws, [("Engine", _Report(findings=(_Finding("OK", "a\x07", "1", ""),)))])

    assert ["OK", "a", "1"] in [row[:3] for row in _cells(ws)]  # sin el carácter de control


def test_connection_report_satisfies_the_structural_view() -> None:
    from util_bbdd.connection import EngineReadinessAdvisor

    report: ReadinessReportView = EngineReadinessAdvisor().assess("mysql", {"pool_size": 2})
    assert ReadinessRows.from_report(report).rows == (
        ("WARNING", "pool_size", "2", "Mínimo 5 para cargas concurrentes"),
    )
