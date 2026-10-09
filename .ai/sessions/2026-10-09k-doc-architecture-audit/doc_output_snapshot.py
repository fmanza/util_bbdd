"""Vuelca a texto la documentación (Word, Excel, HTML) de cada motor del banco para comparar antes/después.

Uso (desde la raíz, con ``UTIL_BBDD_TEST_<MOTOR>_URL`` definidas; ver ``scripts/bench_env.example.sh``)::

    python .ai/sessions/2026-10-09k-doc-architecture-audit/doc_output_snapshot.py build/logs/doc_baseline/before
    python .ai/sessions/2026-10-09k-doc-architecture-audit/doc_output_snapshot.py build/logs/doc_baseline/after
    git diff --no-index build/logs/doc_baseline/before build/logs/doc_baseline/after

Crea y borra el schema de caracterización de ``tests/integration/metadata/characterization/_schemas.py``.
"""

from __future__ import annotations

import html
import os
import re
import sys
import tempfile

from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))

import docx  # noqa: E402
import openpyxl  # noqa: E402

from tests.integration.metadata.characterization._schemas import build, drop, targets  # noqa: E402
from util_bbdd import SyncDatabaseManager  # noqa: E402
from util_bbdd.doc import DatabaseDocumentation  # noqa: E402

_TAG = re.compile(r"<[^>]+>")
_EMBEDDED = re.compile(r"<(style|script)\b.*?</\1>", re.IGNORECASE | re.DOTALL)
_VOLATILE = re.compile(r"\d{4}-\d{2}-\d{2}([ T_]\d{2}:?\d{2}(:?\d{2})?)?|\d{2}/\d{2}/\d{4}( \d{2}:\d{2}(:\d{2})?)?")

IMPRIMIR: dict[str, Any] = {
    "tablas": {k: True for k in ("imprimir_tablas", "ddl", "columnas", "pk", "fk", "unicas", "indices", "charset_collation")},
    "vistas": {"imprimir_vistas": True, "ddl": True, "columnas": True},
    "rutinas": {"imprimir_procedures": True, "imprimir_function": True},
    "almacenamiento": {"imprimir_almacenamiento": True},
    "info_database": {"imprimir_info_database": True, "version": True, "charset_collation": True},
    "columnas_charset": {"charset_collation": True},
    "grafo": {"enabled": True, "max_depth": 2, "direction": "both", "max_subnodes": 2},
}

SCENARIOS: dict[str, tuple[list[str], dict[str, Any]]] = {
    "word_paginated": (["word"], {"images_paginated": True}),
    "word_single": (["word"], {"images_paginated": False}),
    "excel": (["excel"], {}),
    "html": (["html"], {"images_paginated": False}),
    "word_html_paginated": (["word", "html"], {"images_paginated": True}),
}


def _clean(text: str) -> str:
    lines = (" ".join(_VOLATILE.sub("<fecha>", line).split()) for line in text.splitlines())
    return "\n".join(line for line in lines if line) + "\n"


def _word_text(path: Path) -> str:
    document = docx.Document(str(path))
    out: list[str] = [f"[images={len(document.inline_shapes)}]"]
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append("".join(t.text or "" for t in child.iter() if t.tag.endswith("}t")))
        elif tag == "tbl":
            for row in child.iter():
                if row.tag.endswith("}tr"):
                    cells = [
                        "".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t"))
                        for cell in row
                        if cell.tag.endswith("}tc")
                    ]
                    out.append(" | ".join(cells))
    return "\n".join(out)


def _excel_text(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path))
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _html_text(root: Path) -> str:
    parts: list[str] = []
    for page in sorted(root.rglob("*.html")):
        raw = page.read_text(encoding="utf-8")
        parts.append(f"=== {page.relative_to(root).as_posix()}")
        parts.append(html.unescape(_TAG.sub(" ", _EMBEDDED.sub(" ", raw))))
    svgs = sorted(root.rglob("*.svg"))
    parts.append(f"[svg_files={len([s for s in svgs if 'icon-' not in s.name])}]")
    return "\n".join(parts)


def _render(mgr: SyncDatabaseManager, formatos: list[str], graph: dict[str, Any]) -> dict[str, str]:
    out = Path(tempfile.mkdtemp(prefix="doc_snap_"))
    imprimir = {**IMPRIMIR, "grafo": {**IMPRIMIR["grafo"], **graph}}
    DatabaseDocumentation(
        db_manager=mgr, directorio_salida=out, fichero_salida=Path("snap"), imprimir=imprimir, formatos=formatos
    ).generate()
    texts: dict[str, str] = {}
    for docx_file in sorted(out.rglob("*.docx")):
        texts[f"word/{docx_file.name}"] = _word_text(docx_file)
    for xlsx_file in sorted(out.rglob("*.xlsx")):
        texts[f"excel/{xlsx_file.name}"] = _excel_text(xlsx_file)
    html_dirs = sorted({p.parent for p in out.rglob("index.html")})
    for html_dir in html_dirs:
        texts["html"] = _html_text(html_dir)
    return texts


def main() -> None:
    dest = Path(sys.argv[1])
    only = set(sys.argv[2].split(",")) if len(sys.argv) > 2 else None
    dest.mkdir(parents=True, exist_ok=True)
    for target in targets():
        if only and target.engine not in only:
            continue
        if target.url_env and not os.environ.get(target.url_env):
            print(f"skip {target.engine}: {target.url_env} no definida")
            continue
        build(target)
        try:
            db_type, params = target.manager_args()
            for name, (formatos, graph) in SCENARIOS.items():
                with SyncDatabaseManager(db_type, params, use_pool=False) as mgr:
                    texts = _render(mgr, formatos, graph)
                body = "\n".join(f"##### {key}\n{_clean(value)}" for key, value in sorted(texts.items()))
                (dest / f"{target.engine}_{name}.txt").write_text(body, encoding="utf-8")
                print(f"ok {target.engine} {name}")
        finally:
            drop(target)


if __name__ == "__main__":
    main()
