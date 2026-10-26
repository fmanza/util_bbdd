"""Plantilla Excel de la documentación (``imprimir["plantilla_excel"]``; util_bbdd_web G40, 4.1.0).

``ExcelGenerator`` crea el libro desde cero con la paleta de Office en colores RGB fijos y fuentes sin nombre. Con
plantilla, el libro toma el tema de la plantilla, esos colores pasan a ser colores del tema (``accent1``/``accent5``
con su tono: con el tema de Office se ven igual) y las fuentes siguen la fuente del tema; cada hoja copia el
encabezado, el pie y la configuración de página de la primera hoja de la plantilla. Sin portada ni imágenes: openpyxl
no copia hojas entre libros ni las imágenes de un encabezado (``&G``).
"""

from __future__ import annotations

import copy
import re
import zipfile

from pathlib import Path
from typing import Any, ClassVar

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils.exceptions import InvalidFileException

from ...._core.exceptions import DocumentationConfigError


class ExcelTemplate:
    """Tema, fuente y encabezado/pie/página de una plantilla Excel, para aplicarlos a un libro generado.

    Example:
        >>> ExcelTemplate.load(Path("corporativa.xltx")).apply(workbook)
    """

    #: Colores fijos de ``ExcelGenerator`` → (índice de color del tema, tono). En Excel, 4 = ``accent1``, 8 = ``accent5``.
    FIXED_COLORS: ClassVar[dict[str, tuple[int, float]]] = {
        "4472C4": (4, 0.0),  # accent1 de Office 2013+
        "4F81BD": (4, 0.0),  # accent1 de Office 2007
        "366092": (4, -0.25),  # accent1, 25 % más oscuro
        "5B9BD5": (8, 0.0),  # accent5 de Office 2013+
        "2F75B5": (8, -0.25),  # accent5, 25 % más oscuro
        "D9E1F2": (4, 0.8),  # accent1, 80 % más claro (bandas de las filas)
    }

    _THEME_RELS: ClassVar[re.Pattern[bytes]] = re.compile(rb"\br:(?:embed|link|id)\s*=")
    _MINOR_FONT: ClassVar[re.Pattern[bytes]] = re.compile(rb"<a:minorFont>\s*<a:latin\s+typeface=\"([^\"]*)\"")
    _IMAGE_CODE: ClassVar[re.Pattern[str]] = re.compile(r"&G")
    _HEADER_PARTS: ClassVar[tuple[str, ...]] = (
        "oddHeader",
        "oddFooter",
        "evenHeader",
        "evenFooter",
        "firstHeader",
        "firstFooter",
    )

    def __init__(self, theme: bytes | None, sheet: Any) -> None:
        self.theme: bytes | None = theme
        self.sheet: Any = sheet
        match = self._MINOR_FONT.search(theme) if theme else None
        self.minor_font: str | None = match.group(1).decode("utf-8", "replace") if match and match.group(1) else None

    @classmethod
    def load(cls, path: Path) -> ExcelTemplate:
        """Lee la plantilla (``.xlsx`` o ``.xltx``).

        Raises:
            DocumentationConfigError: No es un libro de Excel, no tiene hojas o su tema referencia otros ficheros
                (imágenes): openpyxl no los copiaría y Excel daría el libro por dañado.
        """
        try:
            workbook = load_workbook(path)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
            raise DocumentationConfigError(
                f"La plantilla Excel no es un libro .xlsx/.xltx válido: {path} ({error})", operation="DOCUMENTATION"
            ) from error
        theme: bytes | None = workbook.loaded_theme
        if theme is not None and cls._THEME_RELS.search(theme):
            raise DocumentationConfigError(
                f"El tema de la plantilla Excel {path.name} referencia imágenes u otros ficheros",
                operation="DOCUMENTATION",
            )
        if not workbook.worksheets:
            raise DocumentationConfigError(f"La plantilla Excel {path.name} no tiene hojas", operation="DOCUMENTATION")
        return cls(theme, workbook.worksheets[0])

    @classmethod
    def check(cls, path: Path) -> None:
        """Valida la plantilla sin usarla (lo hace ``DatabaseDocumentation`` al construirse).

        Raises:
            DocumentationConfigError: Ver :meth:`load`.
        """
        cls.load(path)

    def apply(self, workbook: Workbook) -> None:
        """Pone el tema de la plantilla y retoca todas las hojas de ``workbook`` (en memoria, antes de guardarlo)."""
        if self.theme is not None:
            workbook.loaded_theme = self.theme
        for ws in workbook.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    self._recolor(cell)
            self._copy_page_setup(ws)

    def _theme_color(self, color: Any) -> Color | None:
        """Color del tema equivalente a un color fijo conocido (``None`` si no es uno de ellos)."""
        rgb = getattr(color, "rgb", None)
        if getattr(color, "type", None) != "rgb" or not isinstance(rgb, str):
            return None
        mapped = self.FIXED_COLORS.get(rgb[-6:].upper())
        return Color(theme=mapped[0], tint=mapped[1]) if mapped else None

    def _recolor(self, cell: Any) -> None:
        """Rellenos con colores fijos → colores del tema; fuentes sin nombre → la fuente del tema."""
        if isinstance(cell, MergedCell):
            return
        fill = cell.fill
        if fill.fill_type == "solid" and (color := self._theme_color(fill.fgColor)) is not None:
            cell.fill = PatternFill(fill_type="solid", fgColor=color, bgColor=copy.copy(color))
        font = cell.font
        if self.minor_font and font.scheme is None and font.name in (None, "Calibri"):
            new = copy.copy(font)
            new.name, new.scheme = self.minor_font, "minor"
            cell.font = new

    def _copy_page_setup(self, target: Any) -> None:
        """Encabezado, pie y página de la primera hoja de la plantilla, sin imágenes (``&G``)."""
        source = self.sheet
        for part in self._HEADER_PARTS:
            header = copy.deepcopy(getattr(source, part))
            for side in (header.left, header.center, header.right):
                if side.text:
                    side.text = self._IMAGE_CODE.sub("", side.text) or None
            setattr(target, part, header)
        target.HeaderFooter.differentFirst = source.HeaderFooter.differentFirst
        target.HeaderFooter.differentOddEven = source.HeaderFooter.differentOddEven
        for attr in ("orientation", "paperSize", "fitToWidth", "fitToHeight", "scale"):
            setattr(target.page_setup, attr, getattr(source.page_setup, attr))
        target.sheet_properties.pageSetUpPr = copy.copy(source.sheet_properties.pageSetUpPr)
        target.page_margins = copy.copy(source.page_margins)
        target.print_options = copy.copy(source.print_options)
