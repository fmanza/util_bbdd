"""Escritura segura de informes (Excel con openpyxl y Word con python-docx) para toda la librería.

Dos defensas, en un único punto para los informes de ``dump``, ``bulk`` (auditoría), ``comparar`` y ``doc``:

- **Caracteres de control**: openpyxl lanza ``IllegalCharacterError`` (no es ``ValueError`` ni ``OSError``) al asignar una
  cadena con caracteres de control (``\\x00``-``\\x08``, ``\\x0b``, ``\\x0c``, ``\\x0e``-``\\x1f``). Un mensaje de error del
  driver o un comentario del catálogo con uno de ellos abortaba el informe entero. Se eliminan al escribir. python-docx
  falla igual (``ValueError: All strings must be XML compatible``): los mismos caracteres, porque XML 1.0 no los admite.
- **Fórmulas** (H12, audit-code 2026-09-25): openpyxl toma por fórmula toda cadena que empieza por ``=``, así que un
  nombre o comentario del catálogo como ``=HYPERLINK(...)`` se ejecutaría al abrir el fichero. Los informes de la
  librería nunca contienen fórmulas propias: al guardar, toda celda de fórmula se guarda como texto.

Uso:

- Excel: crear el libro con :class:`SafeWorkbook` en lugar de ``openpyxl.Workbook``; el resto del código no cambia. Para
  los libros que crea pandas (``DataFrame.to_excel``), sanear antes con :func:`safe_frame` y aplicar
  :func:`neutralize_formulas` a cada hoja.
- Word: python-docx no tiene un punto único por el que pase todo el texto; los generadores pasan por :func:`safe_text` el
  texto no literal de ``add_heading``, ``add_paragraph``, ``add_run`` y ``celda.text``.
"""

from __future__ import annotations

from collections.abc import Iterable
from typing import Any

import pandas as pd

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.worksheet import Worksheet


def safe_value(value: Any) -> Any:
    """Devuelve ``value`` sin caracteres de control si es una cadena; cualquier otro valor, sin cambios."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def safe_text(value: object) -> str:
    """``str(value)`` sin los caracteres que XML 1.0 no admite (para python-docx). ``None`` pasa a ``""``."""
    if value is None:
        return ""
    cleaned: str = ILLEGAL_CHARACTERS_RE.sub("", value if isinstance(value, str) else str(value))
    return cleaned


def safe_frame(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Copia de ``dataframe`` con :func:`safe_value` aplicado a las columnas de texto (para ``to_excel``)."""
    result = dataframe.copy()
    for column in result.columns:
        if not pd.api.types.is_numeric_dtype(result[column]) and not pd.api.types.is_datetime64_any_dtype(
            result[column]
        ):
            result[column] = result[column].map(safe_value)
    return result


def neutralize_formulas(worksheet: Worksheet) -> None:
    """Guarda como texto las celdas que openpyxl ha tomado por fórmulas."""
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


class SafeWorksheet(Worksheet):
    """``Worksheet`` que elimina los caracteres de control en ``cell``, ``append`` y ``ws["A1"] = valor``.

    No añade estado: :class:`SafeWorkbook` la asigna como clase de las hojas que crea openpyxl.
    """

    # types-openpyxl declara ``cell`` con sobrecargas (``value=None`` → Cell | MergedCell; con valor → Cell): una sola
    # firma con ``Any`` es compatible con todas (config-review 2026-10-02e).
    def cell(self, row: int, column: int, value: Any = None) -> Any:
        return super().cell(row=row, column=column, value=safe_value(value))

    def append(self, iterable: Any) -> None:
        if isinstance(iterable, dict):
            super().append({key: safe_value(val) for key, val in iterable.items()})
        elif isinstance(iterable, Iterable) and not isinstance(iterable, (str, bytes)):
            super().append([safe_value(val) for val in iterable])
        else:
            super().append(iterable)  # openpyxl lanza su propio TypeError

    def __setitem__(self, key: str, value: Any) -> None:
        super().__setitem__(key, safe_value(value))


class SafeWorkbook(Workbook):
    """``openpyxl.Workbook`` cuyas hojas son :class:`SafeWorksheet` y que neutraliza las fórmulas al guardar.

    No admite ``write_only=True``: sus hojas (``WriteOnlyWorksheet``) no pasan por ``cell``/``append`` de ``Worksheet``.
    """

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        if kwargs.get("write_only") or (args and args[0]):
            raise ValueError("SafeWorkbook no admite write_only=True")
        super().__init__(*args, **kwargs)
        for candidate in self.worksheets:  # API pública (``_sheets`` no está en types-openpyxl)
            # ``type(...) is`` excluye las subclases; la variable tipada permite reasignar ``__class__`` (mypy).
            if type(candidate) is Worksheet:
                sheet: Worksheet = candidate
                sheet.__class__ = SafeWorksheet

    # ``title: Any``: types-openpyxl admite ``str | _Decodable | None`` (``_Decodable`` solo existe en los stubs).
    def create_sheet(self, title: Any = None, index: int | None = None) -> Worksheet:
        sheet: Worksheet = super().create_sheet(title=title, index=index)
        if type(sheet) is Worksheet:
            sheet.__class__ = SafeWorksheet
        return sheet

    def save(self, filename: Any) -> None:
        for sheet in self.worksheets:
            if isinstance(sheet, Worksheet):
                neutralize_formulas(sheet)
        super().save(filename)


__all__: list[str] = ["SafeWorkbook", "SafeWorksheet", "neutralize_formulas", "safe_frame", "safe_text", "safe_value"]
