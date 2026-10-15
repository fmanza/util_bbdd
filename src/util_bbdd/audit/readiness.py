"""Hoja «Condiciones del motor» de los Excel de auditoría de ``bulk`` y ``dump`` (``EngineReadinessAdvisor``).

Capa 3 (Regla #6): no importa ``connection``. El informe llega por el contrato estructural :class:`ReadinessReportView`
(lo cumple ``util_bbdd.connection.EngineReadinessReport``) y la hoja se escribe sobre cualquier ``Worksheet`` de
openpyxl: la de ``SafeWorkbook`` (``dump``) o la del ``pd.ExcelWriter`` del resumen de ``bulk``.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import ClassVar, Protocol

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .._core.export_safety import neutralize_formulas, safe_value


class ReadinessFindingView(Protocol):
    """Un hallazgo de la verificación (lo cumple ``util_bbdd.connection.AdvisoryFinding``)."""

    @property
    def level(self) -> str: ...

    @property
    def param(self) -> str: ...

    @property
    def current(self) -> str: ...

    @property
    def recommendation(self) -> str: ...


class ReadinessReportView(Protocol):
    """Informe de la verificación (lo cumple ``util_bbdd.connection.EngineReadinessReport``)."""

    @property
    def db_type(self) -> str: ...

    @property
    def findings(self) -> Sequence[ReadinessFindingView]: ...

    @property
    def checked_at(self) -> datetime: ...

    @property
    def dynamic(self) -> bool: ...

    @property
    def supported(self) -> bool: ...

    @property
    def error(self) -> str | None: ...


#: (título del bloque, informe); ``None`` si no se pudo verificar ese motor.
ReadinessBlock = tuple[str, ReadinessReportView | None]


@dataclass(frozen=True, slots=True)
class ReadinessRows:
    """Contenido de un bloque, igual en todos los Excel: hallazgos ordenados, resumen, verificación y avisos."""

    LEVELS: ClassVar[tuple[str, ...]] = ("CRITICAL", "WARNING", "INFO", "OK")
    NOTICE_UNAVAILABLE: ClassVar[str] = "No se pudieron verificar las condiciones de este motor."
    NOTICE_NO_FINDINGS: ClassVar[str] = "La verificación no ha producido hallazgos."

    rows: tuple[tuple[str, str, str, str], ...]
    summary: tuple[tuple[str, int], ...]
    metadata: tuple[tuple[str, str], ...]
    notices: tuple[str, ...]

    @classmethod
    def from_report(cls, report: ReadinessReportView | None) -> ReadinessRows:
        """Filas del bloque; ``None`` da un bloque solo con el aviso."""
        if report is None:
            return cls(rows=(), summary=(), metadata=(), notices=(cls.NOTICE_UNAVAILABLE,))

        rank = {level: position for position, level in enumerate(cls.LEVELS)}
        rows = tuple(
            sorted(
                ((f.level, f.param, f.current, f.recommendation) for f in report.findings),
                key=lambda row: (rank.get(row[0], len(cls.LEVELS)), row[0], row[1]),
            )
        )
        counts: dict[str, int] = {}
        for row in rows:
            counts[row[0]] = counts.get(row[0], 0) + 1
        summary = tuple((level, counts.get(level, 0)) for level in cls.LEVELS) + tuple(
            (level, count) for level, count in counts.items() if level not in rank
        )

        notices: list[str] = []
        if not report.supported:
            notices.append(f"El motor '{report.db_type}' no tiene verificaciones de condiciones de ejecución.")
        else:
            if not report.dynamic:
                reason = f": {report.error}" if report.error else ""
                notices.append(
                    f"Solo checks estáticos (configuración de conexión), sin variables del servidor{reason}."
                )
            if not rows:
                notices.append(cls.NOTICE_NO_FINDINGS)

        checked_at = report.checked_at.astimezone(UTC) if report.checked_at.tzinfo else report.checked_at
        metadata = (
            ("Motor", report.db_type),
            ("Verificado (UTC)", checked_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Checks dinámicos", "Sí" if report.dynamic else "No"),
        )
        return cls(rows=rows, summary=summary, metadata=metadata, notices=tuple(notices))


class ReadinessSheetWriter:
    """Escribe los bloques de condiciones del motor en una hoja, uno debajo de otro."""

    HEADERS: ClassVar[tuple[str, ...]] = ("Nivel", "Parámetro", "Valor actual", "Recomendación")
    NOTE: ClassVar[str] = (
        "Condiciones verificadas en este proceso por EngineReadinessAdvisor, una vez por engine: si el engine ya había "
        "cargado datos, son las de esa carga."
    )
    _HEADER_FILL: ClassVar[str] = "FF4472C4"
    _LEVEL_FILL: ClassVar[dict[str, str]] = {
        "CRITICAL": "FFFF0000",
        "WARNING": "FFFFC000",
        "INFO": "FFDDEBF7",
        "OK": "FF92D050",
    }

    def write(self, ws: Worksheet, blocks: Sequence[ReadinessBlock]) -> None:
        """Escribe ``blocks`` en ``ws`` y guarda como texto lo que parezca fórmula."""
        title = ws.cell(row=1, column=1, value="Condiciones del motor")
        title.font = Font(size=14, bold=True)
        ws.cell(row=2, column=1, value=self.NOTE).font = Font(italic=True)

        row = 4
        for block_title, report in blocks:
            row = self._write_block(ws, row, block_title, ReadinessRows.from_report(report)) + 1

        for column in range(1, len(self.HEADERS) + 1):
            ws.column_dimensions[get_column_letter(column)].width = (20, 34, 24, 80)[column - 1]
        neutralize_formulas(ws)

    def _write_block(self, ws: Worksheet, row: int, title: str, content: ReadinessRows) -> int:
        """Escribe un bloque desde ``row``; devuelve la fila siguiente."""
        ws.cell(row=row, column=1, value=safe_value(title)).font = Font(size=12, bold=True)
        row += 1
        for notice in content.notices:
            ws.cell(row=row, column=1, value=safe_value(f"⚠️ {notice}")).font = Font(italic=True)
            row += 1

        if content.metadata:
            row = self._write_table(ws, row, ("Propiedad", "Valor"), content.metadata)
        if content.summary:
            row = self._write_table(ws, row, ("Nivel", "Hallazgos"), content.summary, level_column=True)
        if content.rows:
            row = self._write_table(ws, row, self.HEADERS, content.rows, level_column=True)
        return row

    def _write_table(
        self,
        ws: Worksheet,
        row: int,
        headers: Sequence[str],
        rows: Sequence[Sequence[str | int]],
        *,
        level_column: bool = False,
    ) -> int:
        """Cabecera con estilo y filas; la primera columna se colorea por nivel si ``level_column``."""
        for column, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=self._HEADER_FILL)
            cell.alignment = Alignment(horizontal="center")
        row += 1
        for values in rows:
            for column, value in enumerate(values, 1):
                ws.cell(row=row, column=column, value=safe_value(value))
            fill = self._LEVEL_FILL.get(str(values[0])) if level_column else None
            if fill:
                ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=fill)
            row += 1
        return row + 1


__all__: list[str] = [
    "ReadinessBlock",
    "ReadinessFindingView",
    "ReadinessReportView",
    "ReadinessRows",
    "ReadinessSheetWriter",
]
