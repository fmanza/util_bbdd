"""Generador del informe Excel de auditoría para volcados de datos.

Genera un único fichero ``.xlsx`` con tres hojas fijas y tres opcionales, en este orden:
  - ``Auditoría Tablas``: volumetría origen vs. destino por tabla.
  - ``Tablas Faltantes``: tablas del origen sin correspondencia en destino.
  - ``Columnas Inexistentes``: columnas presentes en origen pero ausentes en destino.
  - ``Índices`` (opcional): índices secundarios deshabilitados y reconstruidos.
  - ``Auditoría Bulk`` (opcional): resumen del CSV de auditoría de ``bulk`` y contadores del volcado.
  - ``Secuencias PK`` (opcional): resincronización de los contadores de PK en destino.

Las hojas opcionales se añaden siempre al final, para no mover las demás.
"""

from __future__ import annotations

import logging

from datetime import datetime
from pathlib import Path
from typing import ClassVar

import openpyxl

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .._core.exceptions import DumpReportError
from .._core.export_safety import SafeWorkbook
from ._constraints import ConstraintResult
from ._models import DumpAuditSummary, IndexActionResult, PKSyncResult, TableDumpResult


class DumpReporter:
    """Genera el informe Excel de auditoría de un volcado de datos."""

    logger: ClassVar[logging.Logger] = logging.getLogger(__name__)

    # ------------------------------------------------------------------
    # Paleta de colores (ARGB)
    # ------------------------------------------------------------------

    _COL_HEADER_BLUE: ClassVar[str] = "FF4472C4"
    _COL_HEADER_ORANGE: ClassVar[str] = "FFED7D31"
    _COL_GREEN: ClassVar[str] = "FF92D050"
    _COL_RED: ClassVar[str] = "FFFF0000"
    _COL_YELLOW: ClassVar[str] = "FFFFFF00"
    _COL_ORANGE: ClassVar[str] = "FFFFC000"
    _COL_LIGHT_ORANGE: ClassVar[str] = "FFFCE4D6"
    _COL_WHITE_FONT: ClassVar[str] = "FFFFFFFF"
    _STATUS_BG: ClassVar[dict[str, str]] = {
        "success": "FF92D050",
        "error": "FFFF0000",
        "skipped": "FFFFFF00",
        "missing_target": "FFFFC000",
        "enabled": "FF92D050",
        "failed": "FFFF0000",
        "unverified": "FFFFFF00",
        "rebuilt": "FF92D050",
        "recreated": "FF92D050",
        "unsupported": "FFFFC000",
        "synced": "FF92D050",
        "no_rows": "FF92D050",
    }

    def __init__(self, output_dir: Path) -> None:
        self._output_dir = output_dir

    # ------------------------------------------------------------------
    # Helpers estáticos de escritura de celdas
    # ------------------------------------------------------------------

    @staticmethod
    def _header(ws: Worksheet, row: int, col: int, value: str, color: str = "FF4472C4") -> None:
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True, color=DumpReporter._COL_WHITE_FONT)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _cell(
        ws: Worksheet,
        row: int,
        col: int,
        value: str | int | float | bool | None,
        bg: str | None = None,
    ) -> None:
        cell = ws.cell(row=row, column=col, value=value)
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        if bg:
            cell.fill = PatternFill("solid", fgColor=bg)

    @staticmethod
    def _autowidth(ws: Worksheet) -> None:
        for col_cells in ws.columns:
            width = max((len(str(c.value or "")) for c in col_cells), default=8)
            ws.column_dimensions[get_column_letter(col_cells[0].column or 1)].width = min(width + 4, 65)

    @staticmethod
    def _write_audit_rows(ws: Worksheet, table_results: list[TableDumpResult], cycle_set: set[str]) -> None:
        for row, res in enumerate(table_results, 2):
            bg = DumpReporter._STATUS_BG.get(res.status)
            pos = str(res.load_order_position) if res.load_order_position >= 0 else "—"
            DumpReporter._cell(ws, row, 1, res.table_name)
            DumpReporter._cell(ws, row, 2, pos)
            DumpReporter._cell(ws, row, 3, "Sí" if res.table_name in cycle_set else "No")
            DumpReporter._cell(ws, row, 4, res.source_rows if res.source_rows >= 0 else "—")
            DumpReporter._cell(ws, row, 5, res.target_rows_before if res.target_rows_before >= 0 else "—")
            DumpReporter._cell(ws, row, 6, res.target_rows_after if res.target_rows_after >= 0 else "—")
            DumpReporter._cell(ws, row, 7, res.rows_rejected)
            DumpReporter._cell(ws, row, 8, res.status.upper(), bg=bg)
            DumpReporter._cell(ws, row, 9, ", ".join(res.skipped_columns) if res.skipped_columns else "—")
            DumpReporter._cell(ws, row, 10, res.error_msg or "—")

    @staticmethod
    def _write_constraint_section(
        ws: Worksheet,
        constraint_results: list[ConstraintResult],
        table_count: int,
    ) -> None:
        failed = [r for r in constraint_results if r.status in ("failed", "unverified")]
        if not failed:
            return
        offset = table_count + 4
        DumpReporter._header(
            ws, offset, 1, "Constraints con incidencia al rehabilitar", color=DumpReporter._COL_HEADER_ORANGE
        )
        for col, h in enumerate(["Tabla", "Constraint", "Dialect", "Estado", "Mensaje"], 1):
            DumpReporter._header(ws, offset + 1, col, h, color=DumpReporter._COL_HEADER_ORANGE)
        for i, cr in enumerate(failed, offset + 2):
            DumpReporter._cell(ws, i, 1, cr.table_name)
            DumpReporter._cell(ws, i, 2, cr.constraint_name)
            DumpReporter._cell(ws, i, 3, cr.dialect)
            DumpReporter._cell(ws, i, 4, cr.status.upper(), bg=DumpReporter._STATUS_BG.get(cr.status))
            DumpReporter._cell(ws, i, 5, cr.error_msg or "—")

    def generate(
        self,
        table_results: list[TableDumpResult],
        missing_tables: list[str],
        constraint_results: list[ConstraintResult],
        cycles: list[list[str]],
        index_results: list[IndexActionResult] | None = None,
        audit_summary: DumpAuditSummary | None = None,
        pk_sync_results: list[PKSyncResult] | None = None,
    ) -> Path:
        """Genera el fichero Excel y retorna su ruta.

        Args:
            table_results: Resultado por tabla volcada. ``skipped_columns`` alimenta
                la hoja "Columnas Inexistentes" (columnas presentes en origen y ausentes
                en destino), por lo que el informe no requiere comparar esquemas.
            missing_tables: Tablas del origen que no existen en el destino.
            constraint_results: Resultados de rehabilitación de constraints.
            cycles: Grupos de tablas con FK circular.
            index_results: Resultados de gestión de índices. ``None`` (default) cuando
                ``DumpConfig.manage_indexes=False``: no se añade la hoja "Índices".
            audit_summary: Resumen de la auditoría de ``bulk`` del volcado. ``None`` (default) cuando no hay
                ``DumpConfig.bulk_audit_path`` o no hubo nada que resumir: no se añade la hoja "Auditoría Bulk".
            pk_sync_results: Resultado de resincronizar los contadores de PK (``DumpResult.pk_sync_results``).
                ``None`` (default) cuando ``DumpConfig.sync_pk_sequences=False`` o no hubo tablas que sincronizar:
                no se añade la hoja "Secuencias PK".

        Raises:
            DumpReportError: Ante cualquier fallo al generar el informe: directorio de salida, escritura del fichero
                (disco lleno, permisos, fichero abierto) o un valor que openpyxl no admite. ``output_path`` lleva la
                ruta prevista.
        """
        output_path = self._free_report_path(self._output_dir, datetime.now().strftime("%Y%m%d_%H%M%S"))
        try:
            self._output_dir.mkdir(parents=True, exist_ok=True)

            # SafeWorkbook: quita caracteres de control y guarda como texto lo que parezca fórmula (H12).
            wb = SafeWorkbook()
            if wb.active is not None:
                wb.remove(wb.active)

            self._sheet_audit(wb, table_results, constraint_results, cycles)
            self._sheet_missing_tables(wb, missing_tables)
            self._sheet_missing_columns(wb, table_results)
            if index_results is not None:
                self._sheet_indexes(wb, index_results)
            if audit_summary is not None:
                self._sheet_bulk_audit(wb, audit_summary)
            if pk_sync_results is not None:
                self._sheet_pk_sync(wb, pk_sync_results)

            wb.save(output_path)
        except Exception as exc:
            # No solo OSError: cualquier fallo del informe debe llegar como DumpReportError para que run() conserve el
            # resultado de la carga (IllegalCharacterError, p.ej., no es OSError ni ValueError).
            raise DumpReportError(
                f"Error generando el informe Excel de volcado '{output_path}': {exc}", output_path=output_path
            ) from exc

        self.logger.info("Informe Excel guardado en '%s'.", output_path)
        return output_path

    @staticmethod
    def _free_report_path(output_dir: Path, ts: str) -> Path:
        """``dump_report_<ts>.xlsx``, o con sufijo ``_2``, ``_3``… si ya existe (dos volcados en el mismo segundo).

        Sin reserva atómica: dos procesos que lo comprueben a la vez pueden elegir el mismo nombre.
        """
        candidate = output_dir / f"dump_report_{ts}.xlsx"
        n = 2
        while candidate.exists():
            candidate = output_dir / f"dump_report_{ts}_{n}.xlsx"
            n += 1
        return candidate

    # ------------------------------------------------------------------
    # Hoja 6: Secuencias PK (opcional)
    # ------------------------------------------------------------------

    def _sheet_pk_sync(self, wb: openpyxl.Workbook, pk_sync_results: list[PKSyncResult]) -> None:
        """Hoja "Secuencias PK": contadores de PK resincronizados en destino tras la carga."""
        ws: Worksheet = wb.create_sheet("Secuencias PK")

        headers = ["Tabla", "Columna PK", "Dialect", "MAX(PK)", "Estado", "Mensaje"]
        for col, h in enumerate(headers, 1):
            self._header(ws, 1, col, h)
        ws.freeze_panes = "A2"

        if not pk_sync_results:
            cell = ws.cell(row=2, column=1, value="Sin contadores de PK que sincronizar")
            cell.font = Font(italic=True)
        else:
            for row, pr in enumerate(pk_sync_results, 2):
                self._cell(ws, row, 1, pr.table_name)
                self._cell(ws, row, 2, pr.column_name)
                self._cell(ws, row, 3, pr.engine_dialect)
                self._cell(ws, row, 4, pr.max_value if pr.max_value is not None else "—")
                self._cell(ws, row, 5, pr.status.upper(), bg=self._STATUS_BG.get(pr.status))
                self._cell(ws, row, 6, pr.error_msg or "—")

        self._autowidth(ws)

    # ------------------------------------------------------------------
    # Hoja 5: Auditoría bulk (opcional)
    # ------------------------------------------------------------------

    def _sheet_bulk_audit(self, wb: openpyxl.Workbook, summary: DumpAuditSummary) -> None:
        """Hoja "Auditoría Bulk": métricas del CSV de auditoría y contadores del propio volcado."""
        ws: Worksheet = wb.create_sheet("Auditoría Bulk")
        self._header(ws, 1, 1, "Métrica")
        self._header(ws, 1, 2, "Valor")
        ws.freeze_panes = "A2"

        discrepancia = self._COL_YELLOW if summary.total_filas != summary.audited_rows else None
        rows: list[tuple[str, str | int | float, str | None]] = [
            ("CSV de auditoría", str(summary.audit_path), None),
            ("Operaciones en el CSV", summary.total_operaciones, None),
            ("Filas en el CSV", summary.total_filas, discrepancia),
            ("Filas escritas por el volcado", summary.audited_rows, discrepancia),
            ("Tiempo total (s)", summary.tiempo_total_seg, None),
            ("Tiempo medio por lote (s)", summary.tiempo_promedio_seg, None),
            ("Tablas auditadas", summary.tablas_auditadas, None),
            ("Tabla más activa", summary.tabla_mas_activa or "—", None),
            ("Registros no escritos", summary.audit_failures, self._COL_RED if summary.audit_failures else None),
            (
                "Filas confirmadas sin auditar",
                summary.unaudited_committed_rows,
                self._COL_ORANGE if summary.unaudited_committed_rows else None,
            ),
            (
                "Registros sin datos en destino",
                summary.audit_uncommitted_records,
                self._COL_RED if summary.audit_uncommitted_records else None,
            ),
        ]
        for row, (metrica, valor, bg) in enumerate(rows, 2):
            self._cell(ws, row, 1, metrica, bg=bg)
            self._cell(ws, row, 2, valor, bg=bg)

        nota = ws.cell(
            row=len(rows) + 3,
            column=1,
            value=(
                "«Filas en el CSV» suma el fichero desde el inicio del volcado (incluye otros procesos que escriban en "
                "él); «Filas escritas por el volcado» cuenta solo este volcado."
            ),
        )
        nota.font = Font(italic=True)
        self._autowidth(ws)
        ws.column_dimensions["A"].width = 34  # la nota no debe ensanchar la columna de métricas

    # ------------------------------------------------------------------
    # Hoja 4: Índices gestionados (opcional)
    # ------------------------------------------------------------------

    def _sheet_indexes(
        self,
        wb: openpyxl.Workbook,
        index_results: list[IndexActionResult],
    ) -> None:
        """Hoja "Índices": resultado de deshabilitar/reconstruir índices secundarios."""
        ws: Worksheet = wb.create_sheet("Índices")

        headers = ["Tabla", "Índice", "Dialect", "Acción", "Mensaje"]
        for col, h in enumerate(headers, 1):
            self._header(ws, 1, col, h)
        ws.freeze_panes = "A2"

        if not index_results:
            cell = ws.cell(row=2, column=1, value="Sin índices secundarios gestionados")
            cell.font = Font(italic=True)
        else:
            for row, ir in enumerate(index_results, 2):
                self._cell(ws, row, 1, ir.table_name)
                self._cell(ws, row, 2, ir.index_name)
                self._cell(ws, row, 3, ir.dialect)
                self._cell(ws, row, 4, ir.action.upper(), bg=self._STATUS_BG.get(ir.action))
                self._cell(ws, row, 5, ir.error_msg or "—")

        self._autowidth(ws)

    # ------------------------------------------------------------------
    # Hoja 1: Auditoría de tablas
    # ------------------------------------------------------------------

    def _sheet_audit(
        self,
        wb: openpyxl.Workbook,
        table_results: list[TableDumpResult],
        constraint_results: list[ConstraintResult],
        cycles: list[list[str]],
    ) -> None:
        ws: Worksheet = wb.create_sheet("Auditoría Tablas")

        headers = [
            "Tabla",
            "Posición Carga",
            "FK Circular",
            "Filas Origen",
            "Filas Destino Antes",
            "Filas Destino Después",
            "Filas Rechazadas",
            "Estado",
            "Columnas Omitidas",
            "Error",
        ]
        for col, h in enumerate(headers, 1):
            self._header(ws, 1, col, h)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

        cycle_set = {t for group in cycles for t in group}
        self._write_audit_rows(ws, table_results, cycle_set)
        self._write_constraint_section(ws, constraint_results, len(table_results))
        self._autowidth(ws)

    # ------------------------------------------------------------------
    # Hoja 2: Tablas faltantes en destino
    # ------------------------------------------------------------------

    def _sheet_missing_tables(
        self,
        wb: openpyxl.Workbook,
        missing_tables: list[str],
    ) -> None:
        ws: Worksheet = wb.create_sheet("Tablas Faltantes")

        headers = ["Tabla Origen", "Motivo", "Acción Sugerida"]
        for col, h in enumerate(headers, 1):
            self._header(ws, 1, col, h, color=self._COL_HEADER_ORANGE)
        ws.freeze_panes = "A2"

        if not missing_tables:
            cell = ws.cell(row=2, column=1, value="Sin tablas faltantes en destino")
            cell.font = Font(italic=True)
        else:
            for row, table in enumerate(missing_tables, 2):
                self._cell(ws, row, 1, table, bg=self._COL_LIGHT_ORANGE)
                self._cell(ws, row, 2, "La tabla no existe en la BD destino")
                self._cell(ws, row, 3, "Crear la tabla en destino con el DDL del origen antes de reintentar")

        self._autowidth(ws)

    # ------------------------------------------------------------------
    # Hoja 3: Columnas inexistentes en destino
    # ------------------------------------------------------------------

    def _sheet_missing_columns(
        self,
        wb: openpyxl.Workbook,
        table_results: list[TableDumpResult],
    ) -> None:
        """Hoja "Columnas Inexistentes": columnas presentes en origen y ausentes en destino.

        Se deriva de ``TableDumpResult.skipped_columns`` (calculado durante la carga en
        ``DumpOrchestrator._common_columns``), evitando una comparación de esquemas
        independiente. Las tablas ausentes en destino (``status="missing_target"``) ya
        figuran en la hoja "Tablas Faltantes" y llegan con ``skipped_columns`` vacío.
        """
        ws: Worksheet = wb.create_sheet("Columnas Inexistentes")

        headers = [
            "Tabla",
            "Columna Origen",
            "Valor/Tipo en Origen",
            "Estado en Destino",
            "Acción Sugerida",
        ]
        for col, h in enumerate(headers, 1):
            self._header(ws, 1, col, h, color=self._COL_HEADER_ORANGE)
        ws.freeze_panes = "A2"

        missing_col_pairs = [(res.table_name, col_name) for res in table_results for col_name in res.skipped_columns]

        if not missing_col_pairs:
            cell = ws.cell(row=2, column=1, value="Sin columnas inexistentes en destino")
            cell.font = Font(italic=True)
        else:
            for row, (table_name, col_name) in enumerate(missing_col_pairs, 2):
                self._cell(ws, row, 1, table_name)
                self._cell(ws, row, 2, col_name, bg=self._COL_YELLOW)
                self._cell(ws, row, 3, "exists")
                self._cell(ws, row, 4, "Columna ausente en BD destino")
                self._cell(ws, row, 5, "Columna omitida del volcado — añadir DDL al destino si se requiere")

        self._autowidth(ws)
