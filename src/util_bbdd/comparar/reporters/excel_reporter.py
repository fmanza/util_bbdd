"""Reporter de SchemaDiff en formato Excel usando openpyxl."""

from __future__ import annotations

import logging

from pathlib import Path
from typing import ClassVar

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..models import RoutineDiff, SchemaDiff, TableDiff, ViewDiff
from ._base import STATUS_ES, DdlNormalizer, DiffSummary, build_diff_summary


class ExcelDiffReporter:
    """Genera un reporte Excel del SchemaDiff con hojas por categoría.

    Implementa ``DiffReporterProtocol``.

    Estructura del libro:
    - **Resumen**: tarjeta de metadatos y contadores globales.
    - **Tablas**: fila por tabla con estado y conteo de sub-diferencias.
    - **Columnas**: detalle de diferencias de columna por tabla.
    - **Índices**: detalle de diferencias de índice.
    - **ForeignKeys**: detalle de diferencias de FK.
    - **Vistas**: diferencias de vistas.
    - **Rutinas**: diferencias de rutinas.
    """

    # Paleta de colores por estado (ARGB sin '#')
    _FILL_ADDED: ClassVar[str] = "FFD1FAE5"  # verde claro
    _FILL_REMOVED: ClassVar[str] = "FFFEE2E2"  # rojo claro
    _FILL_MODIFIED: ClassVar[str] = "FFFEF3C7"  # ámbar claro
    _FILL_HEADER: ClassVar[str] = "FF1E3A5F"  # azul oscuro
    _FONT_HEADER: ClassVar[Font] = Font(bold=True, color="FFFFFFFF")
    _ALIGN_CENTER: ClassVar[Alignment] = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _ALIGN_LEFT: ClassVar[Alignment] = Alignment(horizontal="left", vertical="center", wrap_text=True)
    _STATUS_FILL: ClassVar[dict[str, str]] = {
        "added": "FFD1FAE5",
        "removed": "FFFEE2E2",
        "modified": "FFFEF3C7",
    }

    def __init__(self) -> None:
        """Inicializa ExcelDiffReporter con logger."""
        self.logger: logging.Logger = logging.getLogger(__name__)

    def render(self, diff: SchemaDiff, output: Path) -> Path:
        """Genera el fichero ``.xlsx`` en la ruta indicada.

        Args:
            diff: Resultado de la comparación de esquemas.
            output: Ruta completa del fichero ``.xlsx`` a generar.

        Returns:
            La ruta del fichero generado.
        """
        output.parent.mkdir(parents=True, exist_ok=True)
        summary: DiffSummary = build_diff_summary(diff)

        wb = Workbook()
        if wb.active is not None:
            wb.remove(wb.active)

        self._add_summary_sheet(wb, diff, summary)
        if diff.table_diffs:
            self._add_tables_sheet(wb, diff.table_diffs)
            self._add_columns_sheet(wb, diff.table_diffs)
            self._add_indices_sheet(wb, diff.table_diffs)
            self._add_fk_sheet(wb, diff.table_diffs)
        if diff.view_diffs:
            self._add_views_sheet(wb, diff.view_diffs)
        if diff.routine_diffs:
            self._add_routines_sheet(wb, diff.routine_diffs)

        wb.save(output)
        self.logger.info("Excel diff report generado: %s", output)
        return output

    # ------------------------------------------------------------------
    # Helpers de formato
    # ------------------------------------------------------------------

    def _header_row(self, ws: Worksheet, headers: list[str], row: int = 1) -> None:
        """Rellena una fila de cabecera con estilo de encabezado."""
        for col_idx, title in enumerate(headers, 1):
            cell: Cell = ws.cell(row=row, column=col_idx, value=title)
            cell.font = self._FONT_HEADER
            cell.fill = PatternFill("solid", fgColor=self._FILL_HEADER)
            cell.alignment = self._ALIGN_CENTER

    def _get_status_fill(self, status: str) -> PatternFill:
        """Devuelve el relleno de color correspondiente al estado."""
        fg_color: str = self._STATUS_FILL.get(status, "FFFFFFFF")
        return PatternFill("solid", fgColor=fg_color)

    @staticmethod
    def _autofit(ws: Worksheet, min_width: int = 10, max_width: int = 60) -> None:
        """Ajusta el ancho de columnas al contenido."""
        for col in ws.columns:
            max_len: int = max((len(str(cell.value or "")) for cell in col), default=min_width)
            width: int = min(max(max_len + 2, min_width), max_width)
            col_idx: int | None = col[0].column
            if col_idx is not None:
                ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ------------------------------------------------------------------
    # Hojas del libro
    # ------------------------------------------------------------------

    def _add_summary_sheet(self, wb: Workbook, diff: SchemaDiff, summary: object) -> None:
        ws: Worksheet = wb.create_sheet("Resumen")
        ws.freeze_panes = "A2"

        meta: list[tuple[str, str]] = [
            ("Origen", diff.source_label),
            ("Destino", diff.target_label),
            ("Tablas añadidas", str(summary.tables_added)),  # type: ignore[attr-defined]
            ("Tablas eliminadas", str(summary.tables_removed)),  # type: ignore[attr-defined]
            ("Tablas modificadas", str(summary.tables_modified)),  # type: ignore[attr-defined]
            ("Vistas con diff", str(summary.total_views)),  # type: ignore[attr-defined]
            ("Rutinas con diff", str(summary.total_routines)),  # type: ignore[attr-defined]
            ("¿Hay diferencias?", "Sí" if summary.has_differences else "No"),  # type: ignore[attr-defined]
        ]
        if diff.summary:
            meta.append(("Resumen", diff.summary))

        self._header_row(ws, ["Concepto", "Valor"])
        for row_idx, (key, val) in enumerate(meta, 2):
            ws.cell(row=row_idx, column=1, value=key).alignment = self._ALIGN_LEFT
            ws.cell(row=row_idx, column=2, value=val).alignment = self._ALIGN_LEFT

        self._autofit(ws)

    def _add_tables_sheet(self, wb: Workbook, table_diffs: tuple[TableDiff, ...]) -> None:
        ws: Worksheet = wb.create_sheet("Tablas")
        ws.freeze_panes = "A2"
        self._header_row(ws, ["Tabla", "Estado", "Cols diff", "Índices diff", "FKs diff", "Comentarios diff"])

        for row_idx, td in enumerate(table_diffs, 2):
            ws.cell(row=row_idx, column=1, value=td.table_name).alignment = self._ALIGN_LEFT
            estado_cell: Cell = ws.cell(row=row_idx, column=2, value=STATUS_ES.get(td.status, td.status))
            estado_cell.fill = self._get_status_fill(td.status)
            estado_cell.alignment = self._ALIGN_CENTER
            ws.cell(row=row_idx, column=3, value=len(td.column_diffs)).alignment = self._ALIGN_CENTER
            ws.cell(row=row_idx, column=4, value=len(td.index_diffs)).alignment = self._ALIGN_CENTER
            ws.cell(row=row_idx, column=5, value=len(td.fk_diffs)).alignment = self._ALIGN_CENTER
            ws.cell(row=row_idx, column=6, value=len(td.comment_diffs)).alignment = self._ALIGN_CENTER

        self._autofit(ws)

    def _add_columns_sheet(self, wb: Workbook, table_diffs: tuple[TableDiff, ...]) -> None:
        ws: Worksheet = wb.create_sheet("Columnas")
        ws.freeze_panes = "A2"
        self._header_row(ws, ["Tabla", "Estado tabla", "Columna", "Atributo", "Valor origen", "Valor destino"])

        row_idx = 2
        for td in table_diffs:
            for cd in td.column_diffs:
                ws.cell(row=row_idx, column=1, value=td.table_name).alignment = self._ALIGN_LEFT
                estado_cell: Cell = ws.cell(row=row_idx, column=2, value=STATUS_ES.get(td.status, td.status))
                estado_cell.fill = self._get_status_fill(td.status)
                estado_cell.alignment = self._ALIGN_CENTER
                ws.cell(row=row_idx, column=3, value=cd.column_name).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=4, value=cd.attribute).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=5, value=cd.source_value).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=6, value=cd.target_value).alignment = self._ALIGN_LEFT
                row_idx += 1

        self._autofit(ws)

    def _add_indices_sheet(self, wb: Workbook, table_diffs: tuple[TableDiff, ...]) -> None:
        ws: Worksheet = wb.create_sheet("Indices")
        ws.freeze_panes = "A2"
        self._header_row(
            ws,
            ["Tabla", "Índice", "Estado", "Cols origen", "Cols destino", "Único origen", "Único destino"],
        )

        row_idx = 2
        for td in table_diffs:
            for idx in td.index_diffs:
                ws.cell(row=row_idx, column=1, value=td.table_name).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=2, value=idx.index_name).alignment = self._ALIGN_LEFT
                estado_cell: Cell = ws.cell(row=row_idx, column=3, value=STATUS_ES.get(idx.status, idx.status))
                estado_cell.fill = self._get_status_fill(idx.status)
                estado_cell.alignment = self._ALIGN_CENTER
                ws.cell(row=row_idx, column=4, value=idx.source_columns).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=5, value=idx.target_columns).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=6, value="Sí" if idx.source_unique else "No").alignment = self._ALIGN_CENTER
                ws.cell(row=row_idx, column=7, value="Sí" if idx.target_unique else "No").alignment = self._ALIGN_CENTER
                row_idx += 1

        self._autofit(ws)

    def _add_fk_sheet(self, wb: Workbook, table_diffs: tuple[TableDiff, ...]) -> None:
        ws: Worksheet = wb.create_sheet("ForeignKeys")
        ws.freeze_panes = "A2"
        self._header_row(
            ws,
            [
                "Tabla",
                "FK",
                "Estado",
                "Cols origen",
                "Tabla ref. origen",
                "Cols dest.",
                "Tabla ref. destino",
            ],
        )

        row_idx = 2
        for td in table_diffs:
            for fk in td.fk_diffs:
                ws.cell(row=row_idx, column=1, value=td.table_name).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=2, value=fk.fk_name).alignment = self._ALIGN_LEFT
                estado_cell: Cell = ws.cell(row=row_idx, column=3, value=STATUS_ES.get(fk.status, fk.status))
                estado_cell.fill = self._get_status_fill(fk.status)
                estado_cell.alignment = self._ALIGN_CENTER
                ws.cell(row=row_idx, column=4, value=fk.source_columns).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=5, value=fk.source_referred_table).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=6, value=fk.target_columns).alignment = self._ALIGN_LEFT
                ws.cell(row=row_idx, column=7, value=fk.target_referred_table).alignment = self._ALIGN_LEFT
                row_idx += 1

        self._autofit(ws)

    def _add_views_sheet(self, wb: Workbook, view_diffs: tuple[ViewDiff, ...]) -> None:
        ws: Worksheet = wb.create_sheet("Vistas")
        ws.freeze_panes = "A2"
        self._header_row(ws, ["Vista", "Estado", "Definición origen", "Definición destino"])

        for row_idx, vd in enumerate(view_diffs, 2):
            ws.cell(row=row_idx, column=1, value=vd.view_name).alignment = self._ALIGN_LEFT
            estado_cell: Cell = ws.cell(row=row_idx, column=2, value=STATUS_ES.get(vd.status, vd.status))
            estado_cell.fill = self._get_status_fill(vd.status)
            estado_cell.alignment = self._ALIGN_CENTER
            src_ddl = DdlNormalizer.normalize_view_ddl(vd.source_definition) if vd.source_definition else ""
            tgt_ddl = DdlNormalizer.normalize_view_ddl(vd.target_definition) if vd.target_definition else ""
            ws.cell(row=row_idx, column=3, value=src_ddl).alignment = self._ALIGN_LEFT
            ws.cell(row=row_idx, column=4, value=tgt_ddl).alignment = self._ALIGN_LEFT

        self._autofit(ws)

    def _add_routines_sheet(self, wb: Workbook, routine_diffs: tuple[RoutineDiff, ...]) -> None:
        ws: Worksheet = wb.create_sheet("Rutinas")
        ws.freeze_panes = "A2"
        self._header_row(ws, ["Rutina", "Tipo", "Estado", "Definición origen", "Definición destino"])

        for row_idx, rd in enumerate(routine_diffs, 2):
            ws.cell(row=row_idx, column=1, value=rd.routine_name).alignment = self._ALIGN_LEFT
            ws.cell(row=row_idx, column=2, value=rd.routine_type).alignment = self._ALIGN_CENTER
            estado_cell: Cell = ws.cell(row=row_idx, column=3, value=STATUS_ES.get(rd.status, rd.status))
            estado_cell.fill = self._get_status_fill(rd.status)
            estado_cell.alignment = self._ALIGN_CENTER
            src_ddl = DdlNormalizer.normalize_routine_ddl(rd.source_definition) if rd.source_definition else ""
            tgt_ddl = DdlNormalizer.normalize_routine_ddl(rd.target_definition) if rd.target_definition else ""
            ws.cell(row=row_idx, column=4, value=src_ddl).alignment = self._ALIGN_LEFT
            ws.cell(row=row_idx, column=5, value=tgt_ddl).alignment = self._ALIGN_LEFT

        self._autofit(ws)
